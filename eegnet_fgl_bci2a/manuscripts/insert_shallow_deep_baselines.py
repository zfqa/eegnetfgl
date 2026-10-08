"""Insert completed Shallow/Deep ConvNet baselines into the Chinese manuscript."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


DOCX_PATH = (
    Path(__file__).resolve().parent
    / "global_routing_paper"
    / "validation_controlled_global_routing_mi_eeg_manuscript_zh_latest.docx"
)


def replace_paragraph(doc: Document, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            paragraph.text = text
            return
    raise RuntimeError(f"Paragraph not found: {prefix}")


def row_values(row) -> list[str]:
    return [cell.text.strip() for cell in row.cells]


def insert_after(table, match: list[str], values: list[str]) -> None:
    anchor = next((row for row in table.rows if row_values(row)[:3] == match), None)
    if anchor is None:
        raise RuntimeError(f"Table row not found: {match}")
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = value
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(9)
    anchor._tr.addnext(row._tr)


def main() -> None:
    doc = Document(DOCX_PATH)
    baseline_table = doc.tables[9]
    if any("ShallowConvNet" in cell.text for row in baseline_table.rows for cell in row.cells):
        raise RuntimeError("Shallow/Deep rows already exist; refusing duplicate insertion.")

    replace_paragraph(
        doc,
        "对比方法包括 2 s/4 s EEGNetV4",
        "对比方法包括 2 s/4 s EEGNetV4、OAS 黎曼切空间逻辑回归、CSP+收缩 LDA、"
        "ShallowConvNet 和 DeepConvNet，均采用相同的被试内跨会话五随机种子协议。"
        "ShallowConvNet 和 DeepConvNet 在两个数据集、两个时间窗、9 名被试和 5 个随机种子上完成复现。",
    )
    replace_paragraph(
        doc,
        "2a 上，全局路由高于 EEGNetV4",
        "2a 上，全局路由高于 EEGNetV4 2 s（59.31% vs 58.09%）、黎曼 TS-LogReg 2 s（56.94%）、"
        "FBCSP-LDA 2 s（56.38%）、ShallowConvNet 2 s（51.87%）和 DeepConvNet 2 s（48.63%）。"
        "其中，仅相对 CSP-LDA 2 s 的优势在 Holm 校正后仍显著；ShallowConvNet 和 DeepConvNet 的配对检验将在补充材料中报告。"
        "完整时间窗中，黎曼 TS-LogReg 4 s 以 65.25% 最强，FBCSP-LDA 4 s（62.96%）次之，路由并未超过这两个方法。",
    )
    replace_paragraph(
        doc,
        "2b 上，自适应 Riemann KD + spatial loss",
        "2b 上，自适应 Riemann KD + spatial loss 达到 77.12%，与 EEGNetV4 2 s 的 77.14% 基本持平，"
        "并高于黎曼 TS-LogReg 2 s、CSP-LDA 2 s、FBCSP-LDA 2 s、ShallowConvNet 2 s（74.77%）"
        "和 DeepConvNet 2 s（74.87%）。全局路由达到 77.01%，略低于最佳单候选但标准差更小。"
        "DeepConvNet 4 s 达到 77.41%，略高于 2 s 路由，但需要完整 4 s 信号；EEGNetV4 4 s 达到 78.95%，"
        "仍是该数据集最强的已评估完整时间窗神经模型。因此 2b 结果应解释为短窗空间结构引导与验证集选择带来的小幅改善，"
        "而不能宣称超过所有完整时间窗方法。",
    )

    # Insert in reverse order after each anchor so displayed order matches the manuscript.
    for anchor, rows in [
        (["2a", "FBCSP-LDA", "2 s"], [
            ["2a", "DeepConvNet", "2 s", "0.4863", "0.4553", "0.3151"],
            ["2a", "ShallowConvNet", "2 s", "0.5187", "0.5113", "0.3582"],
        ]),
        (["2a", "FBCSP-LDA", "4 s"], [
            ["2a", "DeepConvNet", "4 s", "0.4710", "0.4372", "0.2947"],
            ["2a", "ShallowConvNet", "4 s", "0.5038", "0.4923", "0.3384"],
        ]),
        (["2b", "FBCSP-LDA", "2 s"], [
            ["2b", "DeepConvNet", "2 s", "0.7487", "0.7429", "0.4973"],
            ["2b", "ShallowConvNet", "2 s", "0.7477", "0.7443", "0.4954"],
        ]),
        (["2b", "FBCSP-LDA", "4 s"], [
            ["2b", "DeepConvNet", "4 s", "0.7741", "0.7639", "0.5482"],
            ["2b", "ShallowConvNet", "4 s", "0.7666", "0.7619", "0.5332"],
        ]),
    ]:
        for values in rows:
            insert_after(baseline_table, anchor, values)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
