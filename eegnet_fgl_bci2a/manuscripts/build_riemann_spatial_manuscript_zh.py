from __future__ import annotations

from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_final_manuscript as base
import build_global_routing_manuscript as previous


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "manuscripts" / "riemann_spatial_paper"
FIG_DIR = OUT_DIR / "figures"
DOCX_PATH = OUT_DIR / "riemann_kd_spatial_short_window_mi_eeg_manuscript_zh.docx"
MD_PATH = OUT_DIR / "riemann_kd_spatial_short_window_mi_eeg_manuscript_zh.md"

TITLE = "基于黎曼软标签与空间结构蒸馏的短窗运动想象脑电解码"
SHORT_TITLE = "黎曼引导的短窗 MI-EEG 解码"
BLUE = "2E74B5"
DARK_BLUE = "173A5E"
MID_GRAY = "667085"
BLACK = "1D2939"
LIGHT_BLUE = "E8EEF5"


SUBJECT_2A = [
    (1, 0.6687, 0.7021),
    (2, 0.3090, 0.4139),
    (3, 0.7264, 0.7285),
    (4, 0.5076, 0.5722),
    (5, 0.3229, 0.3722),
    (6, 0.4819, 0.4715),
    (7, 0.6285, 0.6278),
    (8, 0.6833, 0.6986),
    (9, 0.7764, 0.7819),
]

SUBJECT_2B = [
    (1, 0.6825, 0.6875),
    (2, 0.5564, 0.6071),
    (3, 0.5519, 0.6331),
    (4, 0.9556, 0.9563),
    (5, 0.8394, 0.8412),
    (6, 0.8006, 0.7806),
    (7, 0.7850, 0.7931),
    (8, 0.8869, 0.8738),
    (9, 0.8131, 0.7838),
]

SEED_ROWS = [
    (7, 0.601852, 0.593388, 0.469136, 0.777232, 0.776148, 0.554464),
    (42, 0.618827, 0.614865, 0.491770, 0.759524, 0.757286, 0.519048),
    (123, 0.589892, 0.584469, 0.453189, 0.776091, 0.774561, 0.552183),
    (2024, 0.562886, 0.559043, 0.417181, 0.774306, 0.772420, 0.548611),
    (3407, 0.609182, 0.605070, 0.478909, 0.777579, 0.775010, 0.555159),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def set_run_cn(
    run,
    size: float = 11,
    bold: bool = False,
    italic: bool = False,
    color: str = BLACK,
    name: str = "Microsoft YaHei",
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc: Document) -> None:
    """Narrative-proposal preset with a named Chinese-font override."""
    base.configure_document(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = SHORT_TITLE
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_cn(run, size=9, color=MID_GRAY)


def add_title_page(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(105)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_cn(paragraph.add_run(TITLE), size=24, bold=True, color=DARK_BLUE)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(50)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_cn(
        paragraph.add_run("基于 BCI Competition IV-2a/2b 的五随机种子跨会话研究"),
        size=13,
        color=MID_GRAY,
    )

    metadata = [
        ("作者", "[作者姓名]"),
        ("单位", "[院系、机构、城市、国家]"),
        ("通讯作者", "[姓名、邮箱、ORCID]"),
        ("目标期刊", "Medical & Biological Engineering & Computing"),
        ("稿件状态", "中文初稿：Riemann KD + Spatial 主线"),
        ("代码仓库", "https://github.com/zfqa/eegnetfgl"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        set_run_cn(paragraph.add_run(f"{label}："), bold=True, color=DARK_BLUE)
        set_run_cn(paragraph.add_run(value))

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(38)
    paragraph.paragraph_format.left_indent = Inches(0.4)
    paragraph.paragraph_format.right_indent = Inches(0.4)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_cn(
        paragraph.add_run(
            "范围说明：本文的核心模型仅包含一个 4 s 黎曼教师和一个 2 s Compact EEGNet 学生。"
            "教师同时提供类别软标签与协方差空间结构目标，形成单教师、双目标监督框架。"
        ),
        size=10.5,
        italic=True,
        color=MID_GRAY,
    )
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(level=level)
    paragraph.paragraph_format.keep_with_next = True
    size = 16 if level == 1 else 13 if level == 2 else 12
    color = BLUE if level < 3 else DARK_BLUE
    set_run_cn(paragraph.add_run(text), size=size, bold=True, color=color)


def add_body(doc: Document, text: str, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.widow_control = True
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.333
    set_run_cn(paragraph.add_run(text), italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.208
        set_run_cn(paragraph.add_run(item))


def set_table_grid(table, widths: list[float]) -> None:
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")


def add_table(doc: Document, title: str, rows: list[list[str]], widths: list[float]) -> None:
    if abs(sum(widths) - 6.5) > 1e-6:
        raise ValueError(f"Table widths must sum to 6.5 inches: {widths}")
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    set_run_cn(caption.add_run(title), size=9.5, bold=True, color=DARK_BLUE)

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    base.set_table_widths(table, widths)
    set_table_grid(table, widths)
    base.set_repeat_table_header(table.rows[0])

    compact_size = 7.4 if len(rows[0]) >= 7 else 8.1
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            base.set_cell_margins(cell, top=80, start=120, bottom=80, end=120)
            if row_index == 0:
                base.set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index in {0, 1} else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            set_run_cn(
                paragraph.add_run(str(value)),
                size=compact_size,
                bold=row_index == 0,
                color=DARK_BLUE if row_index == 0 else BLACK,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, alt_text: str, width: float = 6.35) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", path.stem)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    set_run_cn(caption_paragraph.add_run(caption), size=9, color=MID_GRAY)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color="#475467") -> None:
    x0, y0 = start
    x1, y1 = end
    draw.line((x0, y0, x1, y1), fill=color, width=5)
    angle = np.arctan2(y1 - y0, x1 - x0)
    for offset in (2.55, -2.55):
        draw.line(
            (
                x1,
                y1,
                x1 + 18 * np.cos(angle + offset),
                y1 + 18 * np.sin(angle + offset),
            ),
            fill=color,
            width=5,
        )


def create_pipeline_figure(path: Path) -> None:
    image = Image.new("RGB", (1800, 850), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "单黎曼教师的双目标监督框架", fill="#173A5E", font=font(42, True))

    def box(x0, y0, x1, y1, fill, title, lines):
        draw.rounded_rectangle((x0, y0, x1, y1), radius=14, fill=fill, outline="#344054", width=3)
        draw.text((x0 + 20, y0 + 18), title, fill="#101828", font=font(27, True))
        y = y0 + 64
        for line in lines:
            draw.text((x0 + 20, y), line, fill="#344054", font=font(22))
            y += 34

    box(55, 165, 340, 360, "#E8F1F8", "同一训练 trial", ["前 2 s：学生输入", "完整 4 s：教师输入", "标签保持一一对应"])
    box(430, 90, 800, 315, "#EAF7F2", "4 s 黎曼教师", ["原始 EEG", "OAS 协方差", "切空间 Logistic 回归", "验证集温度校准"])
    box(430, 440, 800, 665, "#FFF4DE", "2 s Compact EEGNet", ["时间卷积特征", "depthwise 空间卷积", "分类输出"])
    box(900, 90, 1260, 315, "#F2F4F7", "监督目标 1", ["类别概率软标签", "Riemann KD", "权重 0.1"])
    box(900, 440, 1260, 665, "#FCE7F3", "监督目标 2", ["协方差相关结构", "pre-depthwise Gram 约束", "权重 0.2"])
    box(1370, 245, 1740, 510, "#EDE9FE", "训练与推理", ["训练：CE + KD + Spatial", "测试：仅保留 2 s 学生", "无需教师与 4 s 输入"])

    arrow(draw, (340, 245), (430, 200))
    arrow(draw, (340, 300), (430, 545))
    arrow(draw, (800, 200), (900, 200))
    arrow(draw, (800, 545), (900, 545))
    arrow(draw, (1260, 200), (1370, 315))
    arrow(draw, (1260, 545), (1370, 440))
    arrow(draw, (615, 315), (615, 440), "#9D174D")
    draw.text(
        (70, 770),
        "所有教师目标只由训练会话及其验证划分产生；评估会话不参与教师拟合、温度校准或模型选择。",
        fill="#667085",
        font=font(22),
    )
    image.save(path, dpi=(180, 180))


def create_subject_delta_figure(path: Path) -> None:
    image = Image.new("RGB", (1800, 920), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "Riemann KD + Spatial 相对 2 s CE 的逐被试准确率变化", fill="#173A5E", font=font(40, True))
    left, right, top, bottom = 150, 1730, 145, 735
    center = 465
    max_abs = 0.11
    draw.line((left, center, right, center), fill="#667085", width=3)
    group_width = (right - left) / 9
    for subject in range(1, 10):
        x_center = left + (subject - 0.5) * group_width
        values = [
            (SUBJECT_2A[subject - 1][2] - SUBJECT_2A[subject - 1][1], -28, "#2E74B5"),
            (SUBJECT_2B[subject - 1][2] - SUBJECT_2B[subject - 1][1], 28, "#D97706"),
        ]
        for value, offset, color in values:
            y = center - value / max_abs * 300
            y0, y1 = sorted([center, y])
            draw.rectangle((x_center + offset - 22, y0, x_center + offset + 22, y1), fill=color)
        draw.text((x_center - 16, bottom + 20), f"S{subject}", fill="#344054", font=font(21))
    for value in [-0.10, -0.05, 0.05, 0.10]:
        y = center - value / max_abs * 300
        draw.line((left, y, right, y), fill="#E4E7EC", width=2)
        draw.text((35, y - 13), f"{value*100:+.0f} pp", fill="#667085", font=font(18))
    draw.rectangle((1260, 80, 1295, 110), fill="#2E74B5")
    draw.text((1310, 76), "BCI IV-2a", fill="#344054", font=font(21))
    draw.rectangle((1480, 80, 1515, 110), fill="#D97706")
    draw.text((1530, 76), "BCI IV-2b", fill="#344054", font=font(21))
    draw.text(
        (70, 860),
        "正值表示加入黎曼软标签与空间结构监督后准确率提高；每名被试先对 5 个随机种子取均值。",
        fill="#667085",
        font=font(21),
    )
    image.save(path, dpi=(180, 180))


def create_figures() -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        "pipeline": FIG_DIR / "riemann_kd_spatial_pipeline_zh.png",
        "subjects": FIG_DIR / "riemann_kd_spatial_subject_delta_zh.png",
    }
    create_pipeline_figure(figures["pipeline"])
    create_subject_delta_figure(figures["subjects"])
    return figures


def tables() -> dict[str, list[list[str]]]:
    subject_2a = [["被试", "2 s CE", "本文方法", "变化"]]
    subject_2a.extend(
        [f"S{s}", f"{ce:.4f}", f"{method:.4f}", f"{method - ce:+.4f}"]
        for s, ce, method in SUBJECT_2A
    )
    subject_2a.append(["平均", "0.5672", "0.5965", "+0.0293"])

    subject_2b = [["被试", "2 s CE", "本文方法", "变化"]]
    subject_2b.extend(
        [f"S{s}", f"{ce:.4f}", f"{method:.4f}", f"{method - ce:+.4f}"]
        for s, ce, method in SUBJECT_2B
    )
    subject_2b.append(["平均", "0.7635", "0.7729", "+0.0095"])

    seed_table = [["Seed", "2a Acc", "2a F1", "2a Kappa", "2b Acc", "2b F1", "2b Kappa"]]
    seed_table.extend(
        [
            str(seed),
            f"{a_acc:.4f}",
            f"{a_f1:.4f}",
            f"{a_kappa:.4f}",
            f"{b_acc:.4f}",
            f"{b_f1:.4f}",
            f"{b_kappa:.4f}",
        ]
        for seed, a_acc, a_f1, a_kappa, b_acc, b_f1, b_kappa in SEED_ROWS
    )
    seed_table.extend(
        [
            ["平均", "0.5965", "0.5914", "0.4620", "0.7729", "0.7711", "0.5459"],
            ["Seed 均值 SD", "0.0216", "0.0214", "0.0288", "0.0076", "0.0078", "0.0152"],
        ]
    )

    return {
        "datasets": [
            ["数据集", "任务", "通道", "采样率", "跨会话评估"],
            ["BCI IV-2a / BNCI2014_001", "4 分类", "22", "250 Hz", "训练会话用于训练/验证；评估会话仅测试"],
            ["BCI IV-2b / BNCI2014_004", "2 分类", "3", "250 Hz", "训练会话用于训练/验证；评估会话仅测试"],
        ],
        "configuration": [
            ["组成", "论文采用的设置"],
            ["学生", "2 s Compact EEGNet；F1=8，D=2，F2=16，dropout=0.5"],
            ["教师", "4 s Riemann TS-LogReg；原始 EEG -> OAS 协方差 -> 黎曼切空间"],
            ["教师选择", "训练集内部 5 折选择 C；候选 C={0.01, 0.1, 1, 10}"],
            ["教师校准", "训练会话验证集温度校准"],
            ["分类监督", "交叉熵 CE，权重 1.0"],
            ["软标签监督", "Riemann KL，权重 0.1，蒸馏温度 1.0"],
            ["空间监督", "pre-depthwise 通道相关结构 MSE，权重 0.2"],
            ["未启用模块", "额外神经教师、geometry loss、SupCon、diffusion、全局路由"],
            ["测试阶段", "仅使用 2 s Compact EEGNet 学生"],
        ],
        "main": [
            ["数据集", "方法", "Acc", "Macro-F1", "Kappa", "相对 CE"],
            ["2a", "2 s Compact EEGNet CE", "0.5672", "0.5550", "0.4229", "--"],
            ["2a", "Riemann KD + Spatial", "0.5965", "0.5914", "0.4620", "+0.0293"],
            ["2b", "2 s Compact EEGNet CE", "0.7635", "0.7568", "0.5270", "--"],
            ["2b", "Riemann KD + Spatial", "0.7729", "0.7711", "0.5459", "+0.0095"],
        ],
        "seed": seed_table,
        "subject_2a": subject_2a,
        "subject_2b": subject_2b,
        "tests": [
            ["数据集", "比较", "差值", "胜/平/负", "配对 t 检验 p", "Wilcoxon p"],
            ["2a", "本文方法 vs 2 s CE", "+0.0293", "7/0/2", "0.0478", "0.0391"],
            ["2b", "本文方法 vs 2 s CE", "+0.0095", "6/0/3", "0.4434", "0.6523"],
            ["2a", "本文方法 vs 既有路由结果", "+0.0034", "4/0/5", "0.7158", "0.9102"],
            ["2b", "本文方法 vs 既有路由结果", "+0.0026", "5/0/4", "0.8489", "0.8203"],
        ],
        "baselines": [
            ["数据集", "方法", "窗口", "Acc", "Macro-F1", "Kappa"],
            ["2a", "本文方法", "2 s", "0.5965", "0.5914", "0.4620"],
            ["2a", "EEGNetV4 CE", "2 s", "0.5809", "0.5666", "0.4412"],
            ["2a", "Riemann TS-LogReg", "2 s", "0.5694", "0.5570", "0.4259"],
            ["2a", "CSP-LDA", "2 s", "0.5530", "0.5270", "0.4040"],
            ["2a", "Compact EEGNet CE", "4 s", "0.5836", "0.5749", "0.4448"],
            ["2a", "EEGNetV4 CE", "4 s", "0.6025", "0.5894", "0.4700"],
            ["2a", "Riemann TS-LogReg", "4 s", "0.6525", "0.6420", "0.5367"],
            ["2b", "本文方法", "2 s", "0.7729", "0.7711", "0.5459"],
            ["2b", "EEGNetV4 CE", "2 s", "0.7714", "0.7676", "0.5427"],
            ["2b", "Riemann TS-LogReg", "2 s", "0.7159", "0.7080", "0.4317"],
            ["2b", "CSP-LDA", "2 s", "0.7136", "0.7076", "0.4272"],
            ["2b", "Compact EEGNet CE", "4 s", "0.7806", "0.7763", "0.5612"],
            ["2b", "EEGNetV4 CE", "4 s", "0.7895", "0.7852", "0.5790"],
            ["2b", "Riemann TS-LogReg", "4 s", "0.7297", "0.7189", "0.4593"],
        ],
    }


ABSTRACT = (
    "短时间窗可降低运动想象脑机接口的决策延迟，但也会丢失后续事件相关去同步信息。"
    "本文提出一种单黎曼教师的双目标监督方法，将完整 4 s 试次中的类别分布与通道协方差结构迁移到 2 s Compact EEGNet。"
    "教师在原始 4 s EEG 上采用 OAS 协方差、黎曼切空间映射和 Logistic 回归，并通过训练会话验证集进行温度校准。"
    "学生同时优化真实标签交叉熵、权重为 0.1 的黎曼软标签 KL 损失，以及权重为 0.2 的空间相关结构损失；"
    "后者约束 depthwise 空间卷积之前的通道 Gram 结构，使其接近教师协方差导出的相关矩阵。"
    "在 BCI Competition IV-2a 和 2b 上，对 9 名被试执行训练会话到评估会话的被试内测试，并使用 5 个随机种子。"
    "方法在 2a 上取得 59.65% 准确率、59.14% Macro-F1 和 0.4620 Kappa，相对 2 s CE 提高 2.93 个百分点；"
    "被试级配对检验达到未经多重校正的名义显著性。2b 上准确率从 76.35% 提高到 77.29%，但差异不显著。"
    "结果表明，完整时间窗的黎曼几何先验能够改善信息更受限的四分类短窗解码，但其收益依赖任务难度和被试。"
)

KEYWORDS = "运动想象脑电；短时间窗；黎曼几何；知识蒸馏；空间结构；EEGNet"

INTRODUCTION = [
    (
        "运动想象脑电解码需要在识别精度与响应速度之间取得平衡。较长时间窗通常包含更完整的 μ/β 节律变化，"
        "但会增加在线脑机接口的等待时间；较短时间窗可更早输出决策，却更容易受到低信噪比、个体差异和判别信息不足的影响。"
    ),
    (
        "EEGNet 通过时间卷积和 depthwise 空间卷积学习紧凑的时空表示，适合参数受限的脑电场景。"
        "另一方面，黎曼方法把试次协方差矩阵视为对称正定流形上的点，能够直接描述通道之间的二阶关系。"
        "二者具有互补性：EEGNet 擅长端到端特征学习，黎曼模型则提供具有明确几何含义的全试次空间结构。"
    ),
    (
        "现有蒸馏工作多聚焦类别软标签，但仅匹配最终概率并不能保证学生真正学习到教师所依赖的空间关系。"
        "本文因此不再引入第二个神经教师，而是让一个 4 s 黎曼教师同时提供类别级与结构级目标。"
        "这使方法问题更集中：完整时间窗的协方差几何，能否指导 2 s EEGNet 改善空间特征学习？"
    ),
]

CONTRIBUTIONS = [
    "提出一个单黎曼教师、双目标监督的短窗 EEG 框架，将 4 s 类别软标签和协方差空间结构同时迁移到 2 s Compact EEGNet。",
    "将空间损失接入 depthwise 空间卷积之前的通道特征，使学生的空间滤波学习直接受到黎曼相关结构约束。",
    "在 BCI Competition IV-2a 和 2b 上采用 9 被试、5 随机种子、跨会话且无测试泄漏的统一评估。",
    "给出逐被试、跨 seed、长短时间窗和经典基线比较，并如实报告 2a 的主要收益与 2b 的有限收益。",
]


def build_docx(figures: dict[str, Path], table_data: dict[str, list[list[str]]]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    add_heading(doc, "摘要", 1)
    add_body(doc, ABSTRACT)
    paragraph = doc.add_paragraph()
    set_run_cn(paragraph.add_run("关键词："), bold=True, color=DARK_BLUE)
    set_run_cn(paragraph.add_run(KEYWORDS))

    add_heading(doc, "1 引言", 1)
    for text in INTRODUCTION:
        add_body(doc, text)
    add_body(doc, "本文的主要贡献如下：", italic=True)
    add_bullets(doc, CONTRIBUTIONS)

    add_heading(doc, "2 材料与方法", 1)
    add_heading(doc, "2.1 数据集与跨会话协议", 2)
    add_body(
        doc,
        "实验使用 MOABB 提供的 BCI Competition IV-2a（BNCI2014_001）和 BCI Competition IV-2b"
        "（BNCI2014_004）。2a 包含 22 个 EEG 通道和 4 类运动想象任务；2b 包含 3 个 EEG 通道和左右手 2 分类任务。"
        "每名被试独立训练和测试。训练会话用于模型拟合与验证划分，评估会话只用于最终测试。",
    )
    add_table(doc, "表 1 数据集与评估协议", table_data["datasets"], [1.9, 0.7, 0.55, 0.75, 2.6])

    add_heading(doc, "2.2 预处理与泄漏控制", 2)
    add_body(
        doc,
        "MOABB MotorImagery 范式将信号带通至 4-38 Hz，并从提示后 0-4 s 提取完整试次，采样率为 250 Hz。"
        "训练会话按类别分层划分 80% 训练集与 20% 验证集。EEGNet 的逐通道均值和标准差仅在训练集上拟合，"
        "随后应用于验证集和测试集。黎曼教师直接使用未经该神经网络标准化的原始裁剪信号。"
    )
    add_body(
        doc,
        "同一 trial 的前 2 s 输入学生，完整 4 s 输入教师，标签和样本顺序保持一一对应。"
        "测试集不参与教师拟合、C 参数选择、温度校准、标准化统计或早停模型选择。",
    )

    add_heading(doc, "2.3 2 s Compact EEGNet 学生", 2)
    add_body(
        doc,
        "学生采用 Compact EEGNet。第一层时间卷积学习频谱相关滤波器，随后由按时间滤波器分组的 depthwise 卷积"
        "跨全部 EEG 通道学习空间滤波器，再通过可分离卷积和线性分类器输出类别 logits。"
        "输入形状为 batch×1×channels×500，其中 500 对应 2 s。",
    )

    add_heading(doc, "2.4 4 s 黎曼教师", 2)
    add_body(
        doc,
        "对每个完整 4 s 训练 trial X_i 估计 OAS 收缩协方差 C_i。协方差矩阵在仿射不变黎曼度量下映射到切空间，"
        "再由 Logistic 回归完成分类。正则参数 C 从 {0.01, 0.1, 1, 10} 中通过训练集内部 5 折分层交叉验证选择。"
        "教师在训练集上拟合后，使用训练会话验证集进行温度校准，得到每个训练 trial 的类别概率 q_i。",
    )

    add_heading(doc, "2.5 黎曼软标签蒸馏", 2)
    add_body(
        doc,
        "学生输出概率记为 p_i，教师概率记为 q_i。类别级蒸馏损失为 "
        "L_KD = T² KL(q_i(T) || p_i(T))。本实验 T=1.0，权重 λ_KD=0.1。"
        "该项把完整 4 s 协方差分类器的类别间相似性迁移到 2 s 学生，而真实标签仍由交叉熵直接监督。",
    )

    add_heading(doc, "2.6 空间结构蒸馏", 2)
    add_body(
        doc,
        "为避免只学习最终类别概率，本文进一步约束学生的通道关系。对教师协方差 C_i 进行对角标准化，得到相关矩阵 R_i。"
        "学生在时间卷积和批归一化之后、depthwise 空间卷积之前产生特征 H_i，并由 H_i H_i^T 构造通道 Gram 矩阵，"
        "再转换为相关矩阵 S_i。空间损失仅比较矩阵的非对角元素："
        "L_spatial = mean[(offdiag(S_i)-offdiag(R_i))²]。其权重 λ_spatial=0.2。",
    )
    add_body(
        doc,
        "该接入位置具有明确目的：R_i 描述原始通道的协方差关系，因此应在学生完成跨通道 depthwise 混合之前施加约束。"
        "若在分类器末端约束，空间解释会被多层非线性与池化稀释。",
    )

    add_heading(doc, "2.7 总目标与训练设置", 2)
    add_body(
        doc,
        "最终目标为 L = L_CE + 0.1 L_KD + 0.2 L_spatial。实验未启用额外神经教师、geometry loss、"
        "对比学习、扩散增强或全局路由。学生使用 AdamW，学习率 0.001，权重衰减 0.0001，batch size 32，"
        "最多训练 300 epoch，并依据验证集表现早停。随机种子为 7、42、123、2024 和 3407。",
    )
    add_table(doc, "表 2 主实验模型与损失配置", table_data["configuration"], [1.45, 5.05])
    add_figure(
        doc,
        figures["pipeline"],
        "图 1 单个 4 s 黎曼教师向 2 s Compact EEGNet 提供类别软标签和空间相关结构监督。",
        "流程图：同一训练试次的完整四秒信号进入黎曼教师，前两秒信号进入 Compact EEGNet 学生；"
        "教师提供 Riemann KD 与空间结构两个训练目标，测试阶段仅保留两秒学生。",
    )

    add_heading(doc, "2.8 评价指标与统计分析", 2)
    add_body(
        doc,
        "报告准确率、Macro-F1 和 Cohen's Kappa。每个配置在 9 名被试和 5 个随机种子上运行，共 45 个结果。"
        "总体均值先对全部被试和 seed 汇总；逐被试比较先对 5 个 seed 取均值，再执行双侧配对 t 检验与 Wilcoxon"
        "符号秩检验。由于样本量仅为 9 名被试，本文同时报告胜/平/负计数，并避免仅凭单个 p 值作强结论。",
    )

    add_heading(doc, "3 结果", 1)
    add_heading(doc, "3.1 两个数据集的总体结果", 2)
    add_body(
        doc,
        "在四分类 2a 上，本文方法达到 0.5965 准确率、0.5914 Macro-F1 和 0.4620 Kappa，"
        "相对同一 2 s Compact EEGNet CE 基线提高 0.0293。二分类 2b 上准确率由 0.7635 提高至 0.7729，"
        "Macro-F1 与 Kappa 也小幅提高，但增益明显小于 2a。",
    )
    add_table(doc, "表 3 五随机种子主结果", table_data["main"], [0.55, 2.15, 0.85, 1.0, 0.9, 1.05])

    add_heading(doc, "3.2 跨随机种子稳定性", 2)
    add_body(
        doc,
        "2a 的五个 seed 平均准确率范围为 0.5629-0.6188，seed 均值标准差为 0.0216。"
        "2b 的范围为 0.7595-0.7776，标准差仅为 0.0076，说明二分类结果对初始化和训练划分更稳定。"
        "表中的 0.1515 和 0.1115 若被报告，则表示 45 个被试-seed 观测的离散程度，不能误写为 seed 间标准差。",
    )
    add_table(doc, "表 4 每个随机种子的平均结果", table_data["seed"], [0.55, 0.95, 0.95, 0.95, 0.95, 1.05, 1.1])

    add_heading(doc, "3.3 BCI IV-2a 逐被试结果", 2)
    add_body(
        doc,
        "2a 上 9 名被试中有 7 名获得提高。S2 和 S4 的绝对增益最大，分别约为 10.49 和 6.46 个百分点；"
        "S6 和 S7 略有下降。结果说明完整时间窗空间结构对部分低基线被试尤其有帮助，但并非对所有被试稳定有效。",
    )
    add_table(doc, "表 5 BCI IV-2a 逐被试五 seed 平均准确率", table_data["subject_2a"], [0.9, 1.7, 2.0, 1.9])

    add_heading(doc, "3.4 BCI IV-2b 逐被试结果", 2)
    add_body(
        doc,
        "2b 上 6 名被试提高、3 名下降。S2 和 S3 的增益较明显，但 S6、S8 和 S9 出现下降。"
        "由于 2b 仅包含左右手二分类且基线已较高，额外的 4 s 几何监督更容易表现为被试依赖的小幅变化。",
    )
    add_table(doc, "表 6 BCI IV-2b 逐被试五 seed 平均准确率", table_data["subject_2b"], [0.9, 1.7, 2.0, 1.9])
    add_figure(
        doc,
        figures["subjects"],
        "图 2 本文方法相对 2 s CE 的逐被试准确率变化。",
        "双数据集逐被试柱状图。蓝色表示 BCI 2a，橙色表示 BCI 2b；正值为本文方法优于两秒交叉熵基线。",
    )

    add_heading(doc, "3.5 配对统计检验", 2)
    add_body(
        doc,
        "2a 相对 CE 的被试级平均提升达到未经多重比较校正的名义显著性：配对 t 检验 p=0.0478，"
        "Wilcoxon p=0.0391。若将两个数据集或更多基线同时视为同一检验族，该证据在 Holm 校正后不再稳健，"
        "因此本文使用“名义显著”而非“显著优于所有基线”的表述。2b 的差异不显著。",
    )
    add_table(doc, "表 7 被试级配对比较", table_data["tests"], [0.55, 2.15, 0.8, 0.8, 1.05, 1.15])

    add_heading(doc, "3.6 与长短窗基线的比较", 2)
    add_body(
        doc,
        "在 2a 上，本文 2 s 方法高于 2 s EEGNetV4、2 s Riemann TS-LogReg、2 s CSP-LDA 和 4 s Compact EEGNet，"
        "但低于 4 s EEGNetV4 和 4 s Riemann TS-LogReg。在 2b 上，本文方法略高于 2 s EEGNetV4，"
        "但仍低于 4 s Compact EEGNet 与 4 s EEGNetV4。因此，本文贡献是改善紧凑短窗学生，而非宣称全面超过长窗强基线。",
    )
    add_table(doc, "表 8 与已完成基线实验的比较", table_data["baselines"], [0.5, 2.0, 0.65, 0.95, 1.15, 1.25])

    add_heading(doc, "4 讨论", 1)
    add_heading(doc, "4.1 为什么 2a 的收益更明显", 2)
    add_body(
        doc,
        "2a 是四分类、22 通道任务，空间混淆和类别边界更复杂。2 s 学生仅观察试次前半段时，"
        "4 s 协方差所包含的完整通道关系能提供额外归纳偏置。相比之下，2b 只有 3 通道和 2 类，"
        "短窗 CE 已能捕获较多判别信息，额外结构约束的边际收益较小。",
    )

    add_heading(doc, "4.2 软标签与空间结构的互补含义", 2)
    add_body(
        doc,
        "Riemann KD 约束最终类别分布，回答“教师认为该 trial 属于各类别的相对可能性”；"
        "空间损失则约束中间通道关系，回答“学生应当以怎样的空间结构组织早期特征”。"
        "二者来自同一个黎曼教师，但作用层级不同，构成类别级与结构级双目标监督。",
    )

    add_heading(doc, "4.3 实际部署意义", 2)
    add_body(
        doc,
        "黎曼教师和完整 4 s 信号仅在训练期间使用。完成训练后，推理阶段只运行 2 s Compact EEGNet，"
        "不会增加在线采集时长，也不需要实时计算协方差、切空间映射或教师概率。"
        "因此该框架把长窗知识转化为训练期正则化，而保留短窗学生的部署成本。",
    )

    add_heading(doc, "4.4 局限性", 2)
    add_bullets(
        doc,
        [
            "当前统计单位仅为 9 名被试；2a 的 p 值接近 0.05，且在扩大多重比较检验族后不再稳健。",
            "空间损失权重 0.2 与 KL 权重 0.1 仍属于经验设置，需要在嵌套验证或外层交叉验证中确认。",
            "本文尚缺少严格匹配的 CE、仅 Riemann KD、仅 Spatial、Riemann KD + Spatial 四项五 seed 消融表。",
            "2b 上平均收益较小且 3 名被试下降，说明固定空间约束可能需要验证集控制的开启策略。",
            "当前研究为被试内跨会话评估，不能直接推断跨被试迁移能力。",
            "结果尚未在更多公开运动想象数据集或真实在线异步脑机接口中验证。",
        ],
    )

    add_heading(doc, "5 结论", 1)
    add_body(
        doc,
        "本文构建了一个单黎曼教师的短窗 EEG 蒸馏框架。4 s Riemann TS-LogReg 同时向 2 s Compact EEGNet"
        "提供类别软标签和协方差空间结构监督，测试阶段仅保留短窗学生。五随机种子跨会话结果显示，"
        "该方法在 BCI IV-2a 上将准确率从 56.72% 提高到 59.65%，并取得未经多重校正的名义显著性；"
        "在 BCI IV-2b 上从 76.35% 提高到 77.29%，但差异不显著。"
        "这些结果支持“完整时间窗黎曼几何可改善困难四分类短窗解码”的有限结论，"
        "同时表明后续工作应重点完善严格消融、验证集门控和更大规模统计验证。",
    )

    add_heading(doc, "声明", 1)
    declarations = [
        ("基金资助", "[填写基金项目；如无：本研究未获得专项资助。]"),
        ("利益冲突", "作者声明不存在利益冲突。"),
        ("伦理审批", "本研究仅使用公开、去标识化数据集，未新增受试者数据采集。"),
        ("数据可用性", "BCI Competition IV-2a 和 2b 可通过 BNCI Horizon 与 MOABB 获取。"),
        ("代码可用性", "https://github.com/zfqa/eegnetfgl"),
        ("作者贡献", "[按照 CRediT 规范填写。]"),
    ]
    for label, value in declarations:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        set_run_cn(paragraph.add_run(f"{label}："), bold=True, color=DARK_BLUE)
        set_run_cn(paragraph.add_run(value))

    add_heading(doc, "参考文献", 1)
    add_body(doc, "参考文献保留英文题录，以避免期刊名、论文题名和 DOI 翻译造成检索歧义。", italic=True)
    for index, reference in enumerate(previous.REFERENCES, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"{index}. {reference}")
        base.set_run_font(run, name="Arial", size=8.8)

    doc.save(DOCX_PATH)


def markdown_table(rows: list[list[str]]) -> list[str]:
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return lines


def build_markdown(table_data: dict[str, list[list[str]]]) -> None:
    lines = [
        f"# {TITLE}",
        "",
        "## 摘要",
        "",
        ABSTRACT,
        "",
        f"**关键词：** {KEYWORDS}",
        "",
        "## 核心方法",
        "",
        "L = L_CE + 0.1 L_Riemann-KD + 0.2 L_spatial。",
        "",
        "单个 4 s Riemann TS-LogReg 教师同时提供类别软标签和协方差空间结构；测试时仅使用 2 s Compact EEGNet。",
        "",
        "## 主结果",
        "",
    ]
    lines.extend(markdown_table(table_data["main"]))
    lines.extend(["", "## Seed 结果", ""])
    lines.extend(markdown_table(table_data["seed"]))
    lines.extend(["", "## 统计检验", ""])
    lines.extend(markdown_table(table_data["tests"]))
    lines.extend(["", "## 逐被试结果：2a", ""])
    lines.extend(markdown_table(table_data["subject_2a"]))
    lines.extend(["", "## 逐被试结果：2b", ""])
    lines.extend(markdown_table(table_data["subject_2b"]))
    lines.append("")
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    figure_paths = create_figures()
    table_data = tables()
    build_docx(figure_paths, table_data)
    build_markdown(table_data)
    print(DOCX_PATH)
    print(MD_PATH)


if __name__ == "__main__":
    main()
