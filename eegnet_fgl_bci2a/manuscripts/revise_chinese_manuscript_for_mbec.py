from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"C:\Users\10906\Documents\eegfgl 2\eegnet_fgl_bci2a\manuscripts\referenced_chinese\中文论文初稿_参考文献标注版.docx")
OUT_DIR = Path(r"C:\Users\10906\Documents\eegfgl 2\eegnet_fgl_bci2a\manuscripts\referenced_chinese")
OUT = OUT_DIR / "中文论文初稿_投稿前详细修改版.docx"


ABSTRACT = [
    "运动想象脑电（motor imagery EEG, MI-EEG）解码是非侵入式脑机接口的重要任务。多数离线研究采用提示后数秒的完整 EEG 片段进行分类，以获得较稳定的 ERD/ERS 特征；然而，较长分析窗口会增加系统决策延迟，不利于早期反馈和低延迟交互。",
    "本文关注 2 s 短窗口 MI-EEG 解码问题，研究在推理阶段仅使用 2 s EEG 的条件下，训练阶段的 4 s 特权信息能否帮助短窗学生模型保留更多判别能力。为此，本文构建了以 Compact EEGNet 为学生模型的短窗解码框架。在 BCI Competition IV-2a 上，训练阶段引入 4 s Compact EEGNet 教师和 4 s 黎曼教师的软标签；在 BCI Competition IV-2b 上，进一步考察 4 s 黎曼软标签与空间协方差结构监督对 2 s 学生模型的作用。所有教师和协方差目标仅在训练阶段使用，测试阶段仍只部署 2 s 学生模型。",
    "为处理被试间时间模式差异，本文还将固定 2 s、黎曼空间监督、自适应窗口等策略作为候选，并使用训练会话内部验证集进行被试级路由选择。评估会话始终隔离，仅用于最终测试。",
    "在 IV-2a 数据集上，验证集控制路由达到 59.31% 的平均准确率，高于固定 2 s CE 的 56.72% 和固定 2 s 双教师蒸馏的 58.47%，相对固定 CE 的提升在被试配对检验中达到统计显著性。在 IV-2b 数据集上，自适应 Riemann KD 加空间监督达到 77.12%，全局路由达到 77.01%，相对固定 2 s CE 的 76.34% 仅表现为小幅正向趋势，且完整 4 s EEGNet 仍保持更高性能。",
    "这些结果表明，训练期 4 s 特权信息可在四分类 IV-2a 中改善 2 s 短窗解码，但其收益具有数据集和任务依赖性。本文的贡献不是证明 2 s 输入可以普遍替代完整时间窗，而是在严格短窗推理约束下，给出一种可复现的训练期知识迁移和验证集控制策略选择路径。",
]


INTRO_REPLACEMENTS = {
    "intro1": "运动想象脑电（motor imagery electroencephalography, MI-EEG）脑机接口通过识别被试想象肢体运动时产生的脑电模式，实现外部设备控制、康复训练和人机交互等功能[1-2]。现有 MI-EEG 解码方法通常使用数秒级提示后 EEG 信号进行分类，相关分类算法和深度学习方法已在综述研究中得到系统总结[3-6]。较长分析窗口能够积累更充分的事件相关去同步化/同步化（event-related desynchronization/synchronization, ERD/ERS）信息，有利于提高分类稳定性；但它也要求系统在给出决策前等待更多 EEG 信号，从而增加反馈延迟。对于需要早期反馈、连续控制或在线交互的 BCI 应用，降低分析窗口长度本身就是一个重要的工程目标[7-8]。",
    "intro2": "2 s 短窗解码是一种直接的低延迟方案。与完整 4 s 或更长窗口相比，2 s 输入可以更早输出预测结果，但也会丢失 trial 后半段可能存在的判别信息。BCI Competition IV 的公开数据和综述工作表明，运动想象相关模式在 trial 内的出现时间、稳定程度和空间分布具有明显被试差异[9-10]。因此，固定使用经验性的 0-2 s 窗口并不一定适合所有被试：部分被试的有效信息可能在更晚的 2 s 片段中更稳定，另一些被试则可能更适合早期固定窗口。这种时间和被试双重异质性，是短窗 MI-EEG 解码的核心困难。",
    "intro3": "除时间窗选择外，短窗学生模型还无法直接观察完整时间窗中的空间统计结构。黎曼几何方法将 EEG trial 表示为空间协方差矩阵，并在对称正定矩阵流形或切空间中进行分类，已被证明能够有效刻画 MI-EEG 的空间判别结构[11-14]。协方差估计的稳定性会影响这类表征，收缩估计可在小样本条件下提高协方差估计的稳健性[15]。因此，一个自然思路是：在训练阶段允许教师模型使用 4 s 全窗 EEG，提取更完整的类别概率和空间协方差信息；而在推理阶段仍只保留 2 s 学生模型。这样的训练期特权信息设置可以在不增加测试延迟的前提下，为短窗模型提供额外监督。",
    "intro4": "不过，4 s 信息并不必然对每个被试都有益。更长窗口可能包含额外判别信息，也可能引入与当前被试或当前训练目标不匹配的噪声。类似地，自适应 2 s 窗口可改善部分被试的时间对齐，但单独使用检测器选窗并不保证提升最终分类性能。已有共适应校准研究也提示，BCI 系统性能与被试状态和校准策略之间存在相互作用[16]。基于这一点，本文不强制所有被试使用同一种训练策略，而是将固定 2 s CE、固定 2 s 蒸馏、自适应 2 s CE 和自适应蒸馏作为候选路线，由训练会话内部验证集进行被试级选择，评估会话仅用于最终测试。",
    "intro5": "基于上述分析，本文提出一种面向 2 s MI-EEG 解码的训练期特权信息和验证集控制路由框架。本文的核心假设是：在推理阶段严格限制为 2 s 输入时，4 s 神经教师和黎曼几何教师可以为短窗学生提供有用监督；但这种监督的效果具有被试和数据集依赖性，因此需要由验证集控制的策略选择机制限制负迁移。本文并不主张 2 s 解码能够普遍超过完整时间窗解码，而是研究在更早决策约束下如何尽可能保留有效判别性能。本文主要贡献如下：",
}


CONTRIBUTIONS = [
    "（1）提出一种面向 2 s 短窗 MI-EEG 解码的训练期特权信息迁移框架，测试阶段仅使用 2 s EEG 和单个 Compact EEGNet 学生模型。",
    "（2）在 IV-2a 上使用 4 s Compact EEGNet 教师和 4 s 黎曼教师提供软标签监督，在 IV-2b 上进一步评估 4 s 黎曼软标签与空间协方差结构监督，从而分析长窗几何信息对短窗学生的作用边界。",
    "（3）将自适应 2 s 窗口作为候选策略，而不是假设其必然优于固定窗口；通过验证集表现判断其是否适合具体被试。",
    "（4）设计验证集控制的全局路由策略，在测试前从多个 2 s 候选策略中进行被试级选择，以降低固定时间窗或固定监督方式带来的负迁移风险。",
    "（5）在 BCI Competition IV-2a 和 IV-2b 两个公开数据集上，采用 9 名被试、5 个随机种子和跨会话评估协议，与 Compact EEGNet、EEGNetV4、ShallowConvNet、DeepConvNet、CSP-LDA、FBCSP-LDA 和黎曼切空间分类器等基线进行比较[17-20]，并报告 OOF 路由敏感性分析。"
]


METHODS = [
    ("2.1 数据集与跨会话协议", "信号通过 MOABB 获取[21]。本文使用 BCI Competition IV-2a（BNCI2014_001）和 BCI Competition IV-2b（BNCI2014_004）两个公开运动想象 EEG 数据集[9-10]。IV-2a 包含 9 名被试、22 个 EEG 通道和 4 个运动想象类别；IV-2b 包含 9 名被试、3 个 EEG 通道和 2 个类别。所有实验均采用 subject-dependent 的跨会话协议：训练会话用于模型训练和内部验证，评估会话仅用于最终测试。训练会话按类别分层划分为训练子集和验证子集（80/20）。验证集用于早停、温度校准、候选路线选择和路由选择；评估会话不参与任何教师拟合、模型选择、数据标准化参数估计或路由选择。"),
    ("2.2 预处理与时间窗定义", "EEG 信号采用 8-32 Hz 频带滤波，采样率为 250 Hz。每个 trial 从提示后 0-4 s 截取为完整窗口，得到 1000 个采样点；2 s 短窗对应 500 个采样点。固定短窗默认使用 0-2 s。对于需要 4 s 教师的实验，教师输入为同一 trial 的 0-4 s EEG，学生输入为该 trial 对应的 2 s EEG，确保教师软标签和学生样本顺序严格一致。标准化参数仅在训练子集上估计，并应用于验证集和测试集，避免测试集信息泄漏。"),
    ("2.3 学生模型与对比模型", "部署阶段的学生模型为 Compact EEGNet，其输入为 2 s EEG。本文保留 EEGNet 中时间卷积、depthwise 空间卷积和 separable 卷积的轻量结构思想[17, 24]，使用相同的训练协议比较固定 2 s CE、固定 4 s CE 和蒸馏模型。外部深度基线包括 EEGNetV4、ShallowConvNet 和 DeepConvNet[18]；传统基线包括 CSP-LDA、FBCSP-LDA 和黎曼切空间逻辑回归[19-20]。这些基线用于判断所提短窗策略在相同跨会话设置下的相对位置。"),
    ("2.4 训练期教师模型", "IV-2a 中使用两类 4 s 教师。第一类为 4 s Compact EEGNet 教师，使用训练子集的 4 s EEG 和真实标签进行监督训练，并在训练结束后冻结；其输出 logits 经过验证集温度校准后作为学生的神经教师软标签。第二类为 4 s 黎曼教师。对每个 4 s trial 计算 OAS 收缩协方差矩阵[15]，再映射到黎曼切空间，并使用逻辑回归进行分类[11-14]。逻辑回归正则化参数在训练会话内部选择，教师概率同样使用验证集进行温度校准。IV-2b 中主要使用 4 s 黎曼教师，并额外利用其协方差结构构造空间监督目标。"),
    ("2.5 蒸馏损失与空间结构监督", "对于普通 CE 训练，学生模型仅优化交叉熵损失。对于双教师蒸馏，学生同时学习真实标签和两个教师的软标签；该设计来源于知识蒸馏思想，即用教师输出的软概率约束学生模型[22-23]。损失写为 L = CE(y, p_s) + alpha_eeg KL(p_s^T || p_eeg^T) + alpha_rie KL(p_s^T || p_rie^T)，其中 p_s 为学生概率，p_eeg 和 p_rie 分别为神经教师和黎曼教师概率，T 为温度。对于 IV-2b 的 Riemann KD + 空间监督，损失写为 L = CE + lambda_KD KL(p_s^T || p_rie^T) + lambda_spatial L_spatial。空间损失约束学生 depthwise 空间卷积前后的空间相关结构，使学生中间表示与由 4 s 协方差得到的空间目标保持一致。所有辅助损失只在训练阶段使用，测试阶段不调用教师模型或协方差分类器。"),
    ("2.6 自适应短窗与验证集控制路由", "自适应窗口模块首先在训练会话内部训练一个 MI-active 检测器，并在 0-2 s 可检测范围内以 0.5 s 检测窗口和 0.25 s 步长扫描 EEG 片段。检测器输出用于确定候选 2 s 分类窗口的起点。由于单独的检测式选窗可能与最终分类最优窗口不一致，本文不将自适应窗口作为默认最优策略，而是把固定 CE、固定蒸馏、自适应 CE 和自适应蒸馏共同作为候选。全局路由在每个被试和随机种子下，仅根据训练会话内部验证集表现选择候选策略，并将该策略应用到隔离的评估会话。"),
    ("2.7 评价指标与统计分析", "本文报告 accuracy、macro-F1 和 Cohen's kappa。主要结果以 5 个随机种子下的 9 被试均值表示，标准差按随机种子层面计算；当表格采用被试层面标准差时，在表注中明确说明。IV-2a 的主要统计比较采用被试配对 t 检验和 Wilcoxon signed-rank 检验，对比验证集控制路由与固定 2 s CE。对于仍未完成配对检验的消融矩阵，本文仅作描述性解释，不作显著性声称。")
]


RESULTS_REPLACEMENTS = {
    "r1": "验证集控制路由在 IV-2a 上达到 59.31% 的准确率，而固定 2 s CE 为 56.72%（提升 +2.59 个百分点）。在种子平均后，9 名被试中有 8 名被试获得提升，且差异具有统计显著性（配对 t 检验 p=0.0055；Wilcoxon p=0.0078）。固定 2 s 双教师蒸馏达到 58.47%，低于路由结果但高于固定 CE；路由与固定双教师之间的差异不显著，说明主要优势来自被试特异性的候选选择，而不是某一个候选策略在所有被试上均占优。需要注意的是，单独的自适应 CE 在 IV-2a 上低于固定 CE，提示检测式窗口选择和最终分类最优窗口之间并不完全一致。",
    "r2": "在 IV-2b 上，自适应 Riemann KD 加空间监督达到 77.12%，全局路由达到 77.01%，固定 CE 为 76.34%。该增益小于 IV-2a，和 IV-2b 较少的通道数以及二分类任务设置相一致。完整 4 s Compact EEGNet CE 达到 78.06%，仍高于所有 2 s 策略，说明 IV-2b 上短窗方法主要提供低延迟条件下的小幅性能保持，而不是超过完整时间窗。匹配辅助损失的消融显示，中等联合权重（0.10+0.10）取得最高均值；但该消融目前仍为描述性分析，因为尚未针对最终空间监督矩阵重新计算与 CE 的配对检验。",
    "r3": "本文将所提出 2 s 策略与复现的 ShallowConvNet、DeepConvNet、经典方法、黎曼方法和 EEGNetV4 基线进行比较。除 EEGNet 与 Shallow/DeepConvNet 外，近年来还出现了 FBCNet、EEG Conformer、多视图 CNN、注意力时序卷积网络、胶囊网络和 TSception 等 EEG 深度解码模型[25-30]。在 IV-2a 上，验证集控制路由超过复现的 ShallowConvNet、DeepConvNet、FBCSP-LDA、Riemann TS-LogReg 和 EEGNetV4 2 s 基线。在 IV-2b 上，自适应 Riemann KD 加空间监督接近 EEGNetV4 2 s CE，并超过 Shallow/Deep 2 s 基线；但其均值未超过 EEGNetV4，因此不应表述为全面优于 EEGNetV4。",
    "r4": "4 s 模型仍是强参考。在 IV-2a 上，4 s Riemann TS-LogReg 和 FBCSP-LDA 明显高于所提 2 s 路由；在 IV-2b 上，4 s EEGNetV4 和 4 s Compact EEGNet 也高于 2 s 策略。因此，本文的短窗贡献应理解为在更早决策约束下保留和改善 2 s 解码性能，而不是超过所有完整时间窗方法。",
    "r5": "IV-2b 辅助损失消融为描述性分析：Riemann KD + 空间监督在 0.10+0.10 权重下达到最高被试级均值（accuracy 0.7648，macro-F1 0.7617，kappa 0.5296）。这些数值是在 5 种子平均后对被试进行汇总得到的，因此其标准差反映被试异质性，不能直接与表 2 中的种子级标准差比较。仅 Riemann KD 和仅空间损失均取得接近结果，说明两种监督信号可能存在互补性，但当前证据不足以支持强显著性结论。"
}


DISCUSSION = [
    ("4 讨论", "本文的主要发现是，在推理阶段不携带 4 s 教师的条件下，训练期特权信息和验证集控制策略选择能够改善 IV-2a 的 2 s 短窗解码。该结果支持短窗 MI-EEG 解码的一种实用路径：完整时间窗不必直接进入部署模型，而可以在训练阶段作为教师信息，为低延迟学生模型提供额外监督。"),
    (None, "IV-2a 与 IV-2b 的差异说明，该策略的收益具有明显任务依赖性。IV-2a 是 22 通道四分类任务，空间协方差结构和类别关系更丰富，因此 4 s 神经教师和黎曼教师可能提供了更多短窗学生难以直接观察的信息。相比之下，IV-2b 只有 3 个通道且为二分类任务，固定 2 s CE 已经达到较高水平，留给蒸馏和空间监督的改进空间较小。因此，IV-2b 结果应被解释为跨数据集边界验证，而不是强性能突破。"),
    (None, "自适应窗口实验进一步揭示了一个重要现象：MI-active 检测并不等价于分类最优窗口选择。检测器可以在验证集上识别运动想象相关活动，但由检测器确定的窗口在 IV-2a 上并未稳定优于固定 0-2 s CE。这说明短窗分类依赖的不只是活动是否出现，还包括该时间片段内类别间空间模式是否足够可分。因此，本文将自适应窗口作为路由候选，而不是把它作为默认最优策略。"),
    (None, "路由选择的稳定性是本文最重要的局限。虽然主实验中的评估会话始终隔离，但每个随机种子只使用一个内部验证划分进行候选选择。补充 OOF 路由分析显示，IV-2a 上 OOF routing 的准确率为 57.06% ± 0.94%，仅略高于固定 2 s CE，并低于单验证划分路由结果。这提示路由策略对验证划分存在敏感性。未来工作应采用内部 K 折或 OOF 路由作为默认选择机制，并报告验证集排名与测试集排名的一致性。"),
    (None, "本文还有三个限制。第一，所有实验均为被试依赖训练，尚未证明跨被试迁移能力。第二，实验采用公开离线数据集，不能直接证明异步在线 BCI 中的实时表现。第三，部分消融结果仍为描述性分析，尤其是 IV-2b 的 Riemann KD 与空间损失矩阵，还需要完整配对检验和多重比较校正。尽管如此，当前结果已经说明，在短窗约束下引入训练期 4 s 特权信息是可行的，并且其有效性需要结合数据集复杂度和被试差异来判断[3, 5, 14]。"),
]


CONCLUSION = "本文提出并验证了一种面向 2 s MI-EEG 短窗解码的训练期特权信息和验证集控制路由框架。结果表明，4 s 神经教师和黎曼教师能够在 IV-2a 四分类任务中提升 2 s 学生模型性能，而 Riemann KD 加空间监督在 IV-2b 二分类任务中仅带来小幅正向趋势。完整 4 s 模型仍是强参考，因此本文的意义在于低延迟约束下的性能保持和受控提升，而不是用 2 s 输入全面替代完整时间窗。未来工作需要进一步验证 OOF/K 折路由、跨被试迁移和在线异步场景。"


def style_run(run, east_asia="宋体", western="Times New Roman", size=10.5) -> None:
    run.font.name = western
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)


def set_text(paragraph, text: str, size=10.5) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    style_run(run, size=size)


def insert_paragraph_after(paragraph, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.addnext(new_para._p)
    # Put it exactly after original.
    paragraph._p.addnext(new_para._p)
    if style is not None:
        new_para.style = style
    if text:
        run = new_para.add_run(text)
        style_run(run)
    return new_para


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def find_paragraph(doc: Document, exact: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"Paragraph not found: {exact}")


def replace_by_prefix(doc: Document, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_text(paragraph, text)
            return
    raise ValueError(f"Prefix not found: {prefix}")


def remove_between(doc: Document, start_heading: str, end_heading: str) -> None:
    start = find_paragraph(doc, start_heading)
    end = find_paragraph(doc, end_heading)
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def insert_section_after_heading(doc: Document, heading: str, sections: list[tuple[str, str]]) -> None:
    base = find_paragraph(doc, heading)
    current = base
    for subheading, body in sections:
        h = insert_paragraph_after(current, subheading, style="Heading 2")
        current = h
        p = insert_paragraph_after(current, body)
        current = p


def set_table_note(table, note: str) -> None:
    p = table._tbl.getnext()
    while p is not None and p.tag.endswith("p"):
        # Stop if this is a non-empty paragraph that starts the next section/caption.
        break
    # Use simple insertion after table XML.
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    para = table._parent.add_paragraph()
    para._p.getparent().remove(para._p)
    new_p.addnext(para._p)
    table._tbl.addnext(para._p)
    run = para.add_run(note)
    style_run(run, size=9)
    para.paragraph_format.space_after = Pt(6)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(SRC)

    # Abstract: replace five original abstract paragraphs with a bounded version.
    for idx, text in zip(range(3, 8), ABSTRACT):
        set_text(doc.paragraphs[idx], text)
    # Remove the old broad abstract closing sentence; the new abstract already contains a bounded conclusion.
    delete_paragraph(doc.paragraphs[8])

    replace_by_prefix(doc, "运动想象脑电（motor imagery electroencephalography", INTRO_REPLACEMENTS["intro1"])
    replace_by_prefix(doc, "2 s 短窗解码是一种直接的早期决策方案", INTRO_REPLACEMENTS["intro2"])
    replace_by_prefix(doc, "除时间窗选择外", INTRO_REPLACEMENTS["intro3"])
    replace_by_prefix(doc, "但是同时在研究中我发现", INTRO_REPLACEMENTS["intro4"])
    replace_by_prefix(doc, "基于上述分析", INTRO_REPLACEMENTS["intro5"])
    for old_prefix, new_text in zip(["（1）", "（2）", "（3）", "（4）", "（5）"], CONTRIBUTIONS):
        replace_by_prefix(doc, old_prefix, new_text)

    # Replace Methods body between heading and Results, preserving Table 1 and Fig. 1 objects in-place is fragile;
    # the manuscript has tables/images as inline objects after captions, so we keep the table/figure captions and add
    # detailed method text before the existing Table 1 caption.
    replace_by_prefix(doc, "信号通过 MOABB 获取", METHODS[0][1])
    # Turn old 2.2 into the new second methods subsection.
    replace_by_prefix(doc, "部署阶段的学生模型", METHODS[2][1])
    old_h = find_paragraph(doc, "2.2 训练期特权信息与路由")
    set_text(old_h, "2.3 学生模型与对比模型")
    try:
        old_h.style = "Heading 2"
    except Exception:
        pass

    table_caption = find_paragraph(doc, "表 1｜数据集和跨会话评估协议。")
    # Add preprocessing after table caption and before the existing table by inserting before caption would be harder;
    # place it after table caption text but before Methods 2.3 if layout order is acceptable in Word.
    p = table_caption
    for heading, body in METHODS[1:2]:
        h = insert_paragraph_after(p, heading, style="Heading 2")
        p = insert_paragraph_after(h, body)

    # Add method subsections after the current 2.3 paragraph and before Figure 1 caption.
    fig1 = find_paragraph(doc, "图 1｜训练期特权信息与部署阶段短窗推理。4 s 教师和协方差目标仅在训练阶段可用。部署阶段使用 2 s 学生模型；若选中自适应候选，则额外使用可选检测器。")
    anchor = doc.paragraphs[[p._p for p in doc.paragraphs].index(fig1._p) - 1]
    current = anchor
    for heading, body in METHODS[3:]:
        h = insert_paragraph_after(current, heading, style="Heading 2")
        current = insert_paragraph_after(h, body)

    replace_by_prefix(doc, "验证集控制路由在 IV-2a", RESULTS_REPLACEMENTS["r1"])
    replace_by_prefix(doc, "在 IV-2b 上，自适应 Riemann", RESULTS_REPLACEMENTS["r2"])
    replace_by_prefix(doc, "本文将所提出 2 s 策略", RESULTS_REPLACEMENTS["r3"])
    replace_by_prefix(doc, "4 s 模型仍是强参考", RESULTS_REPLACEMENTS["r4"])
    replace_by_prefix(doc, "IV-2b 辅助损失消融为描述性分析", RESULTS_REPLACEMENTS["r5"])

    # Replace Discussion section content.
    set_text(find_paragraph(doc, "4 讨论"), DISCUSSION[0][0])
    replace_by_prefix(doc, "本文主要结果是", DISCUSSION[0][1])
    replace_by_prefix(doc, "主要局限是路由选择的稳定性", DISCUSSION[1][1])
    anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("IV-2a 与 IV-2b 的差异说明"):
            anchor = paragraph
            break
    if anchor is None:
        anchor = find_paragraph(doc, DISCUSSION[1][1])
    current = anchor
    for _, text in DISCUSSION[2:]:
        current = insert_paragraph_after(current, text)

    replace_by_prefix(doc, "训练期特权信息和验证集控制", CONCLUSION)

    # Add table footnotes clarifying std definitions.
    notes = [
        "注：所有划分均在训练会话内部完成，评估会话仅用于最终测试。",
        "注：除另有说明外，数值为 9 名被试在 5 个随机种子下的种子级均值；± 表示 across-seed 标准差。",
        "注：带 ± 的深度基线为 5 个随机种子的种子级均值；未标注 ± 的条目表示当前稿件尚未整理对应 seed-level 标准差，应在正式投稿前补齐或移至补充材料。",
        "注：4 s 结果仅作为完整时间窗参考，不与 2 s 短窗方法进行同延迟比较。",
        "注：本表为被试级描述性消融，± 表示 5 种子平均后的跨被试标准差，不能直接与表 2 的 seed-level 标准差比较。"
    ]
    # Appending notes after tables can perturb figure order, so only add notes if no identical note exists.
    existing = "\n".join(p.text for p in doc.paragraphs)
    for table, note in zip(doc.tables, notes):
        if note not in existing:
            set_table_note(table, note)

    # Basic formatting pass.
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            if run.text:
                style_run(run, size=10.5 if not paragraph.text.strip().startswith("注：") else 9)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
