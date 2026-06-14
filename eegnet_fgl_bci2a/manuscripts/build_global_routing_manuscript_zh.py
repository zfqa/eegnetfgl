from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_final_manuscript as base
import build_global_routing_manuscript as english


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "manuscripts" / "global_routing_paper"
FIG_DIR = OUT_DIR / "figures_zh"
DOCX_PATH = OUT_DIR / "validation_controlled_global_routing_mi_eeg_manuscript_zh.docx"
MD_PATH = OUT_DIR / "validation_controlled_global_routing_mi_eeg_manuscript_zh.md"

ROUTING_ROOT = ROOT / "results" / "global_routing_multiseed"
WEIGHT_ROOT = ROOT / "results" / "weight_matched_kd" / "analysis"
PUBLICATION_ROOT = ROOT / "results" / "publication_complete_analysis"

TITLE = "面向短窗运动想象脑电解码的神经-黎曼双教师验证集控制全局路由"
SHORT_TITLE = "短窗 MI-EEG 的验证集控制路由"

ABSTRACT = (
    "缩短运动想象脑电（MI-EEG）的分析窗口有助于减少分类前所需的信号长度，但固定时间窗和单一监督方式"
    "无法稳定适用于所有被试。本文提出验证集控制全局路由框架，在四个 2 s Compact EEGNet 候选策略中进行"
    "被试级选择：固定窗交叉熵、检测器选择的自适应窗交叉熵、固定窗双教师知识蒸馏，以及自适应窗双教师"
    "知识蒸馏。双教师分别为冻结的 4 s 神经 EEGNet 和基于 4 s 协方差切空间逻辑回归的黎曼教师。候选策略"
    "仅依据训练会话中的验证集准确率选择，评估会话始终隔离。本文在 BCI Competition IV-2a 和 2b 的 9 名"
    "被试上采用 5 个随机种子开展跨会话实验。2a 上，全局路由将固定 2 s 交叉熵的准确率由 56.72% 提升至"
    "59.31%（配对 t 检验 p=0.0055），但跨 seed 共识选择仅达到 57.11%，相对固定策略不显著，说明单次"
    "验证划分存在选择不稳定性。2b 上，路由由 76.34% 提升至 77.37%，但差异不显著。等总权重消融显示，"
    "2a 上双教师 0.2+0.2 达到 58.52%，显著高于黎曼单教师 0.4（p=0.0192），相对神经单教师 0.4 接近"
    "显著（p=0.0567）；2b 上各蒸馏策略差异均不显著。推理阶段不调用教师，在 RTX 3070 Ti 上，固定策略"
    "每 trial 约需 1.01-1.42 ms，自适应策略约需 2.36-3.21 ms。结果支持双教师对四分类短窗解码的互补"
    "作用，但验证集路由仍需更严格的折外选择协议。本文方法应被解释为一种短窗策略选择与训练期知识迁移"
    "框架，而非对完整时间窗解码的普遍替代。"
)

KEYWORDS = "脑机接口；脑电；运动想象；时间路由；知识蒸馏；黎曼几何；EEGNet"

INTRODUCTION = [
    (
        "运动想象脑机接口通过脑电图推断想象运动，可用于交流、康复和设备控制。许多离线解码研究使用提示后"
        "数秒的 EEG。较长时间窗提供更多时序证据，但也会延长系统在作出决策前必须等待的信号时长。因此，"
        "短窗解码对响应式系统具有吸引力，不过早期片段可能仅包含不完整的事件相关去同步或同步模式。"
    ),
    (
        "紧凑卷积网络与基于协方差的黎曼分类器为 MI-EEG 提供了互补视角。EEGNet 从波形中直接学习时间滤波"
        "和深度可分离空间滤波；黎曼方法则以对称正定通道协方差矩阵表示每个 trial，并使用流形距离或切空间"
        "特征进行分类。部署时希望模型足够紧凑，而训练阶段的完整时间窗协方差模型可提供较稳健的几何信息。"
    ),
    (
        "知识蒸馏可用具有额外信息的教师监督受限学生。本研究关注“时间特权”场景：教师在训练时观察 4 s 信号，"
        "学生在训练和推理时仅观察 2 s。初步实验显示，神经教师、黎曼教师与检测器选择的时间窗均具有明显的"
        "被试依赖性；对部分被试有益的策略也会对另一些被试造成负迁移。"
    ),
    (
        "因此，本文将短窗解码表述为验证集控制的策略路由问题。对每个被试和随机种子分别训练固定窗 CE、"
        "自适应窗 CE、固定窗双教师 KD、自适应窗双教师 KD 四个完整候选模型，并只依据训练会话验证准确率"
        "选择一个候选。这里的“全局”表示每个被试和训练种子只选择一条路线，而不是对每个测试 trial 使用"
        "混合专家门控。测试时不调用教师，仅当选择自适应路线时使用时间检测器。"
    ),
]

CONTRIBUTIONS = [
    "提出同时适配时间窗与长窗教师知识使用方式的验证集控制全局路由框架。",
    "构建由 4 s 神经 EEGNet 与 4 s 黎曼切空间逻辑回归共同监督 2 s Compact EEGNet 的双教师训练方案。",
    "在 BCI Competition IV-2a/2b 上完成 9 被试、5 随机种子、跨会话且严格隔离测试集的实验矩阵。",
    "补充等总 KD 权重消融、跨 seed 路由稳健性、推理复杂度和逐类别误差分析，区分教师互补、路由误选与短窗信息上限。",
]

METHODS = {
    "2.1 数据集与实验协议": [
        (
            "BCI Competition IV-2a 通过 MOABB 的 BNCI2014_001 获取，包含 9 名被试、22 个 EEG 通道和"
            "左手、右手、双脚、舌部 4 类运动想象。2b 通过 BNCI2014_004 获取，包含 9 名被试、3 个双极"
            "EEG 通道和左手/右手二分类。两套数据均以 250 Hz 采样。"
        ),
        (
            "实验采用被试内、跨会话协议。指定训练会话按类别分层划分为 80% 训练集和 20% 验证集，评估会话"
            "保持为完全独立的测试集。划分使用随机种子 42、2024、3407、7 和 123。模型训练、教师拟合、"
            "检测器训练、超参数选择、checkpoint 选择和路由选择均只使用训练会话。"
        ),
    ],
    "2.2 预处理与数据泄漏控制": [
        (
            "EEG trial 采用 8-32 Hz 带通滤波，并截取提示后的 0-4 s。逐通道标准化参数仅在训练子集上拟合，"
            "再原样应用于验证集与测试集。固定短窗使用 0-2 s；自适应窗长度始终为 2 s，但起点由检测器选择。"
        ),
        (
            "双教师训练中，学生的 2 s 输入和教师的 4 s 输入均从同一个 trial 裁剪并在批处理前保持对应关系。"
            "标签和样本顺序不被打乱。教师只产生训练集和验证集软标签，测试阶段仅调用 2 s 学生模型。测试"
            "会话不参与时间窗、阈值、超参数、模型 checkpoint 或路线选择。"
        ),
    ],
    "2.3 Compact 学生与长窗教师": [
        (
            "所有路由候选分类器均采用 PyTorch Compact EEGNet，输入形状为 [batch, 1, channels, samples]。"
            "模型使用 F1=8、深度倍增因子 D=2、F2=16、64 样本时间卷积核、ELU、平均池化和 dropout=0.5。"
            "2a 的 4 s 神经教师为 Compact EEGNet；2b 的神经教师为 EEGNetV4 风格模型，采用精确 same padding、"
            "Keras 风格 BatchNorm、125 样本时间核、16 样本可分离卷积核以及 max-norm 约束。教师均在同一"
            "被试的训练子集上训练并冻结。"
        ),
        (
            "黎曼教师使用 Oracle Approximating Shrinkage（OAS）估计 4 s trial 协方差矩阵，按照仿射不变"
            "黎曼度量投影到黎曼均值处的切空间，再使用多项逻辑回归分类。正则参数 C 从 "
            "{0.01, 0.03, 0.1, 0.3, 1, 3, 10} 中依据验证集选择。最终配置不进行全局黎曼对齐。"
        ),
        (
            "双教师目标函数为 L=CE+alpha_E*KL_E+alpha_R*KL_R，其中 CE 为真实标签交叉熵，KL_E 和 KL_R"
            "分别传递 4 s 神经教师和黎曼教师的类别分布。主实验设置 alpha_E=0.4、alpha_R=0.4、T=1.0。"
        ),
    ],
    "2.4 自适应时间窗候选": [
        (
            "2a 使用类别感知 MI-active 检测器，以 0.5 s 窗口和 0.25 s 步长扫描。早期窗口标记为非 MI，"
            "稳定 MI 窗口使用 trial 类别，过渡窗口不参与训练。检测阈值和连续阳性窗口数量由训练会话验证"
            "数据确定，并在五随机种子路由实验前固定。"
        ),
        (
            "2b 使用 EEGNetV4 检测器，以 1 s 窗口和 0.25 s 步长扫描提示后 0-2 s。置信度在相邻 3 个"
            "窗口上平滑，选择峰值置信度对应的起点并向后截取 2 s。两套检测器反映数据集在通道数和任务结构"
            "上的差异，本文不将其宣称为统一的起始点检测算法。"
        ),
    ],
    "2.5 验证集控制全局路由": [
        (
            "每个被试和随机种子均训练固定 CE、自适应 CE、固定双教师 KD、自适应双教师 KD 四个候选。所有"
            "候选在同一验证集上评估，选择准确率最高者。完全相同时采用保守优先级：固定 CE、固定双教师、"
            "自适应 CE、自适应双教师。选择完成后，才在评估会话上测试所选路线。"
        ),
        (
            "固定路线在部署时只运行 Compact EEGNet；自适应路线运行时间检测器后再运行 Compact EEGNet。"
            "无论选择哪条路线，两个教师均仅在训练期使用。"
        ),
    ],
    "2.6 优化、指标与统计检验": [
        (
            "神经候选最多训练 300 epoch，batch size=32，学习率 0.001，权重衰减 0.0001，早停 patience=100。"
            "报告准确率、macro-F1 和 Cohen's kappa。先对每个被试的 5 个随机种子取均值，再在 9 个被试均值"
            "上进行配对 t 检验和 Wilcoxon 符号秩检验。涉及多条外部基线时，在每个数据集内进行 Holm 校正。"
        ),
        (
            "对比方法包括 2 s/4 s EEGNetV4、OAS 黎曼切空间逻辑回归和 CSP+收缩 LDA，均采用相同的被试内"
            "跨会话五随机种子协议。ShallowConvNet 和 DeepConvNet 尚未纳入完成的实验矩阵。"
        ),
    ],
    "2.7 等总 KD 权重消融": [
        (
            "为区分双教师互补与总蒸馏权重增大造成的影响，新增两组等权对照。总 KD 权重为 0.4 时，比较"
            "神经单教师 0.4、黎曼单教师 0.4 和双教师 0.2+0.2；总 KD 权重为 0.8 时，比较神经单教师 0.8、"
            "黎曼单教师 0.8 和双教师 0.4+0.4。每个配置均在两个数据集的 9 名被试和 5 个随机种子上运行。"
        )
    ],
    "2.8 路由稳健性、复杂度与类别分析": [
        (
            "为评估路线选择对单次验证划分的敏感性，进行了留一 seed 的跨 seed 共识分析：对每个被试和待评估"
            "seed，仅使用其余 4 个 seed 的验证集平均准确率选择路线，再读取被留出 seed 的测试结果。该分析"
            "不使用测试集选路，但由于不同 seed 的训练子集存在重叠，应视为稳健性诊断而非严格的 trial 级"
            "嵌套交叉验证。"
        ),
        (
            "推理复杂度统计学生和检测器的可训练参数、卷积/全连接 MACs，并在 CPU 与 RTX 3070 Ti 上测量"
            "batch=1 的每 trial 前向延迟。类别分析使用 45 次运行汇总的混淆计数计算 precision、recall 和 F1。"
        ),
    ],
}

RESULTS = {
    "3.1 BCI Competition IV-2a 主结果": [
        (
            "全局路由达到 59.31% 准确率、58.42% macro-F1 和 0.4575 kappa；固定 2 s CE 分别为 56.72%、"
            "55.50% 和 0.4229。准确率提高 2.59 个百分点，9 名被试中 8 名提高、1 名下降。配对 t 检验"
            "p=0.0055，Wilcoxon p=0.0078；若仅对两个数据集的主比较进行 Holm 校正，2a 仍显著。"
        ),
        (
            "固定双教师候选达到 58.47%。路由进一步提高 0.84 个百分点，但路由与固定双教师的差异不显著"
            "（p=0.1741）。因此，路由收益主要来自利用被试差异并避免统一采用较差候选，而不是相对最强单一"
            "候选产生大幅提升。"
        ),
    ],
    "3.2 BCI Competition IV-2b 跨数据集结果": [
        (
            "2b 上，全局路由达到 77.37% 准确率、76.94% macro-F1 和 0.5473 kappa；固定 2 s CE 为"
            "76.34%、75.66% 和 0.5269。准确率提高 1.02 个百分点，9 名被试中 7 名提高，但未达到显著性"
            "（配对 t 检验 p=0.0720，Wilcoxon p=0.0703）。"
        ),
        (
            "自适应双教师在不经过路线选择时的平均准确率为 77.44%，略高于路由后的 77.37%。这说明有限验证"
            "数据上的最佳路线并不一定最大化测试均值，2b 仅支持方向一致的外部验证，而不能证明稳定增益。"
        ),
    ],
    "3.3 神经教师与黎曼教师基础消融": [
        (
            "2a 上，CE、神经单教师 0.4、黎曼单教师 0.4、双教师 0.4+0.4 的准确率分别为 56.72%、57.42%、"
            "56.55% 和 58.47%。双教师相对 CE 提高 1.75 个百分点（t 检验 p=0.0122；Wilcoxon p=0.0156），"
            "而两个单教师相对 CE 均不显著。"
        ),
        (
            "2b 上，四种配置分别达到 76.35%、76.67%、76.33% 和 76.77%。双教师相对 CE 仅提高 0.42 个"
            "百分点且不显著（p=0.4574）。较少类别和仅 3 个通道可能降低了附加教师知识的边际价值。"
        ),
    ],
    "3.4 等总 KD 权重消融": [
        (
            "等权对照进一步检验了教师互补性。2a 上，总权重 0.4 时，双教师 0.2+0.2 达到 58.52%，高于"
            "神经单教师 0.4 的 57.42% 和黎曼单教师 0.4 的 56.55%。双教师相对黎曼单教师差异显著"
            "（p=0.0192，Wilcoxon p=0.0273），相对神经单教师接近显著（p=0.0567）。"
        ),
        (
            "总权重 0.8 时，2a 的双教师 0.4+0.4 为 58.47%，神经单教师 0.8 为 58.03%，黎曼单教师 0.8"
            "为 56.89%，但两项配对差异均不显著。双教师 0.2+0.2 与 0.4+0.4 几乎相同（58.52% vs 58.47%，"
            "p=0.8919），表明简单增大总 KD 权重没有额外收益。"
        ),
        (
            "2b 上，双教师 0.2+0.2 为 76.81%，双教师 0.4+0.4 为 76.77%，神经单教师 0.4/0.8 为"
            "76.67%/76.69%，黎曼单教师 0.4/0.8 为 76.33%/75.52%。所有等权双教师与单教师比较均不显著。"
        ),
    ],
    "3.5 与神经及经典基线的比较": [
        (
            "2a 上，全局路由高于 EEGNetV4 2 s（59.31% vs 58.09%）、黎曼 TS-LogReg 2 s（56.94%）和"
            "CSP-LDA 2 s（55.30%）。其中，仅相对 CSP-LDA 2 s 的优势在 Holm 校正后仍显著。完整时间窗中，"
            "黎曼 TS-LogReg 4 s 以 65.25% 最强，路由并未超过该方法。"
        ),
        (
            "2b 上，全局路由略高于 EEGNetV4 2 s（77.37% vs 77.14%），但差异不显著。EEGNetV4 4 s 达到"
            "78.95%，是该数据集最强的已评估方法。路由的贡献因此应定位为短窗策略选择，而非不限制时间窗时的"
            "最高分类准确率。"
        ),
        (
            "路由策略的平均有效信号终点为 2a 的 2.11 s 和 2b 的 2.20 s，明显短于 4 s 参考，但该指标只"
            "描述所需 EEG 样本终点，不等同于在线系统的完整决策延迟。"
        ),
    ],
    "3.6 路线使用与被试异质性": [
        (
            "每个数据集共 45 次路由运行。2a 中固定路线占 64.4%，2b 中占 62.2%；双教师路线分别占 48.9%"
            "和 46.7%。没有一个候选策略占据绝对优势，说明时间窗与教师效用均依赖被试和随机划分。"
        ),
        (
            "仅用于诊断的测试集 oracle 分析显示：即使事后选择最佳 2 s 候选，2a 的 S1、S2、S3 和 2b 的"
            "S9 仍低于 4 s CE，提示短窗信息上限；2b 的 S1、S7、S8 存在可匹配或超过 4 s 的 2 s 候选，"
            "但验证集未稳定选中，提示路线选择误差。oracle 结果未用于正式选模或主结果。"
        ),
    ],
    "3.7 跨 seed 路由稳健性": [
        (
            "留一 seed 共识路由在 2a 上达到 57.11%，固定 CE 为 56.72%，差值仅 0.39 个百分点且不显著"
            "（p=0.7127；Wilcoxon p=0.7422）。共识路由还显著低于原单次验证划分路由的 59.31%"
            "（p=0.0107）。"
        ),
        (
            "2b 上，共识路由为 77.04%，固定 CE 为 76.34%，提高 0.69 个百分点但不显著（p=0.1787）；"
            "与原路由 77.37% 的差异也不显著。该结果表明当前路由选择在 2a 尤其依赖单次验证划分，正式结论"
            "应强调验证控制的探索性收益，并将 trial 级嵌套交叉验证作为后续必要验证。"
        ),
    ],
    "3.8 推理复杂度与逐类别表现": [
        (
            "固定 2 s Compact EEGNet 在 2a/2b 上分别包含 2,420/1,634 个执行参数和 5.89/0.86 MMAC。"
            "加入检测器后分别为 3,526/3,500 个参数和 11.25/4.83 MMAC。RTX 3070 Ti 上固定路线延迟约为"
            "1.01/1.42 ms，自适应路线约为 2.36/3.21 ms；教师不参与推理。"
        ),
        (
            "2a 中双脚类别最难：CE 的 F1 为 0.468，双教师提高至 0.492。左手、右手和舌部 F1 也分别由"
            "0.593、0.598、0.595 提高至 0.610、0.626、0.599。2b 中双教师主要改善右手 recall"
            "（0.697 提高至 0.719），左右手 F1 分别为 0.782 和 0.758。"
        ),
    ],
}

DISCUSSION = [
    (
        "本文的主要结果是：验证集控制路由在四分类 2a 上显著提高了 Compact EEGNet 的短窗解码性能，但在"
        "二分类 2b 上仅获得较小且不显著的增益。该结果并非由单个普遍最优的专家产生，而是来自候选策略的"
        "被试依赖性以及验证集选择。"
    ),
    (
        "等总权重消融使双教师机制的解释更严格。2a 上，双教师 0.2+0.2 在总 KD 权重与单教师 0.4 完全一致"
        "时仍达到最高均值，并显著高于黎曼单教师、接近显著高于神经单教师。这排除了“仅因双教师总权重更大”"
        "这一简单解释，并支持神经决策边界与协方差几何信息具有一定互补性。不过，相对神经单教师的证据尚未"
        "越过 0.05 阈值，因此不宜将教师互补表述为已被完全证实。"
    ),
    (
        "2b 的等权结果显示，固定 2 s 基线已较高，神经单教师接近双教师，而提高黎曼教师权重反而下降。这说明"
        "三通道二分类任务中协方差教师的附加信息有限，并且教师权重需要根据数据集与被试校准。"
    ),
    (
        "跨 seed 共识分析揭示了比平均准确率更关键的问题：原路由在 2a 上的显著增益未能在更保守的路线选择"
        "方式下保持。当前路由比较四个候选，却只依赖一个 20% 验证划分，容易放大验证噪声。因此，论文可以"
        "报告原路由的显著结果，但必须同时报告稳健性诊断，并避免将其解释为已稳定泛化的最终路由器。"
    ),
    (
        "与 4 s 模型的比较支持谨慎的实际意义。2a 路由数值上略高于 4 s Compact EEGNet，2b 略低，均未"
        "显著。证据更适合支持“使用更短信号片段保持接近性能”，而不是“2 s 普遍优于 4 s”。若要强化该主张，"
        "应预先设定非劣效界值并开展正式非劣效检验。"
    ),
    (
        "复杂度结果表明训练期教师不会增加部署负担。固定路线只有约 1 ms 的 GPU 前向延迟，自适应路线也低于"
        "4 ms。实际在线延迟仍由 2 s EEG 采集时间主导，但模型计算不是主要瓶颈。"
    ),
    (
        "类别分析显示，2a 中双教师对最困难的双脚类别和右手类别改善较明显，这与黎曼协方差对传感运动空间"
        "关系的敏感性相符。由于这里只分析汇总混淆计数，未来仍需对每个被试的类别改变量与教师置信度进行"
        "配对统计。"
    ),
    (
        "下一步最重要的改进不是继续增加候选，而是在训练会话内部使用严格的 K 折折外预测估计候选效用，并"
        "设置安全阈值：只有当自适应或教师路线在多个内部折上稳定优于固定 CE 时才启用，否则回退固定路线。"
    ),
]

LIMITATIONS = [
    "路由器在分别训练的候选之间进行被试级选择，不是端到端或 trial 级混合专家模型。",
    "主路由结果使用单个 20% 验证划分；跨 seed 共识分析显示其稳定性有限。",
    "2a 与 2b 使用不同的检测器目标和结构，降低了跨数据集方法的一致性。",
    "跨 seed 共识仍不是严格的 trial 级嵌套交叉验证，因为不同 seed 的训练子集存在重叠。",
    "显著提升主要出现在 2a；2b 的提升较小且不显著。",
    "方法未显著超过固定双教师或完整时间窗强基线，不能宣称不受时间窗限制的 SOTA。",
    "当前实验为被试内训练，不能直接推断跨被试迁移或无校准部署能力。",
    "尚未完成 ShallowConvNet、DeepConvNet 等卷积基线以及真实异步在线 BCI 测试。",
]

CONCLUSION = (
    "本文将短窗 MI-EEG 解码表述为验证集控制的时间窗与监督策略路由问题。9 被试、5 随机种子实验表明，"
    "原全局路由在 BCI IV-2a 上显著优于固定 2 s Compact EEGNet，在 2b 上获得方向一致但不显著的提高。"
    "等总权重消融进一步表明，2a 的双教师优势不能仅由总 KD 权重增大解释，并支持神经教师与黎曼教师存在"
    "一定互补性。与此同时，跨 seed 共识路由未显著超过固定 CE，说明当前路线选择仍受验证划分影响。推理"
    "阶段仅保留学生和可选检测器，计算延迟较低。综合而言，该方法为短窗训练期知识迁移提供了有价值的证据，"
    "但在形成稳定的通用路由框架之前，还需要严格的折外选择、统一检测器和前瞻性在线验证。"
)


def cn_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def load_data() -> dict[str, pd.DataFrame]:
    return {
        "runs": pd.read_csv(ROUTING_ROOT / "all_subject_seed_results.csv"),
        "subject_means": pd.read_csv(ROUTING_ROOT / "subject_mean_across_seeds.csv"),
        "selection": pd.read_csv(ROUTING_ROOT / "selection_counts.csv"),
        "consensus": pd.read_csv(
            ROUTING_ROOT / "cross_seed_consensus" / "cross_seed_consensus_subject_summary.csv"
        ),
        "complexity": pd.read_csv(PUBLICATION_ROOT / "inference_complexity.csv"),
        "classes": pd.read_csv(PUBLICATION_ROOT / "per_class_confusion_metrics.csv"),
        "weight_summary": pd.read_csv(WEIGHT_ROOT / "weight_matched_summary.csv"),
        "weight_tests": pd.read_csv(WEIGHT_ROOT / "weight_matched_paired_tests.csv"),
    }


def create_pipeline_figure(path: Path) -> None:
    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "验证集控制全局路由", fill="#173A5E", font=cn_font(42, True))

    def box(x0, y0, x1, y1, fill, title, lines):
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=fill, outline="#344054", width=3)
        draw.text((x0 + 22, y0 + 18), title, fill="#101828", font=cn_font(28, True))
        y = y0 + 66
        for line in lines:
            draw.text((x0 + 22, y), line, fill="#344054", font=cn_font(23))
            y += 36

    box(60, 150, 350, 350, "#E8F1F8", "训练 EEG", ["同一 trial 的 4 s 信号", "仅训练/验证使用"])
    box(430, 105, 790, 245, "#EAF7F2", "固定窗候选", ["0-2 s CE", "0-2 s 双教师 KD"])
    box(430, 300, 790, 470, "#FFF4DE", "自适应窗候选", ["检测器选择 2 s 片段", "CE 或双教师 KD"])
    box(875, 180, 1215, 390, "#F2F4F7", "验证集路由器", ["比较四个候选", "选择验证准确率最高者", "相同时优先固定 CE"])
    box(1300, 180, 1720, 390, "#EDE9FE", "评估会话测试", ["仅运行所选 2 s 策略", "推理阶段不使用教师", "自适应路线才用检测器"])
    box(430, 540, 790, 690, "#FCE7F3", "训练期教师", ["4 s 神经 EEGNet", "4 s 黎曼 TS-LogReg"])

    def arrow(x0, y0, x1, y1, color="#475467"):
        draw.line((x0, y0, x1, y1), fill=color, width=5)
        angle = np.arctan2(y1 - y0, x1 - x0)
        for offset in (2.55, -2.55):
            draw.line(
                (x1, y1, x1 + 18 * np.cos(angle + offset), y1 + 18 * np.sin(angle + offset)),
                fill=color,
                width=5,
            )

    arrow(350, 220, 430, 175)
    arrow(350, 285, 430, 370)
    arrow(790, 175, 875, 245)
    arrow(790, 385, 875, 330)
    arrow(1215, 285, 1300, 285)
    arrow(610, 540, 610, 470, "#9D174D")
    draw.text((65, 610), "全局：每个被试/seed 选择一条路线", fill="#667085", font=cn_font(21))
    image.save(path, dpi=(180, 180))


def create_subject_figure(data: dict[str, pd.DataFrame], path: Path) -> None:
    frame = data["subject_means"]
    image = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 35), "被试平均准确率变化：全局路由相对固定 2 s CE", fill="#173A5E", font=cn_font(40, True))
    left, right, top, bottom = 150, 1730, 140, 760
    center = (top + bottom) // 2
    max_abs = 0.08
    draw.line((left, center, right, center), fill="#667085", width=3)
    group_width = (right - left) / 9
    for subject in range(1, 10):
        x_center = left + (subject - 0.5) * group_width
        for offset, dataset, color in [(-28, "bci2a", "#2E74B5"), (28, "bci2b", "#D97706")]:
            row = frame[(frame["dataset"] == dataset) & (frame["subject"] == subject)].iloc[0]
            value = float(row["global_test_acc"] - row["fixed_ce_test_acc"])
            y = center - value / max_abs * ((bottom - top) / 2)
            y0, y1 = sorted([center, y])
            draw.rectangle((x_center + offset - 23, y0, x_center + offset + 23, y1), fill=color)
        draw.text((x_center - 15, bottom + 24), f"S{subject}", fill="#344054", font=cn_font(21))
    for value in [-0.08, -0.04, 0.04, 0.08]:
        y = center - value / max_abs * ((bottom - top) / 2)
        draw.line((left, y, right, y), fill="#E4E7EC", width=2)
        draw.text((35, y - 13), f"{value*100:+.0f} pp", fill="#667085", font=cn_font(18))
    draw.rectangle((1260, 70, 1295, 100), fill="#2E74B5")
    draw.text((1310, 67), "BCI IV-2a", fill="#344054", font=cn_font(21))
    draw.rectangle((1480, 70, 1515, 100), fill="#D97706")
    draw.text((1530, 67), "BCI IV-2b", fill="#344054", font=cn_font(21))
    draw.text((80, 835), "正值表示验证集控制路由的准确率更高。", fill="#667085", font=cn_font(22))
    image.save(path, dpi=(180, 180))


def create_selection_figure(data: dict[str, pd.DataFrame], path: Path) -> None:
    selection = data["selection"]
    routes = ["fixed_ce", "fixed_dual", "adaptive_ce", "adaptive_dual"]
    labels = ["固定 CE", "固定双教师", "自适应 CE", "自适应双教师"]
    colors = ["#2E74B5", "#16A085", "#D97706", "#9B51E0"]
    image = Image.new("RGB", (1600, 820), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "45 次被试-seed 运行中的路线选择频次", fill="#173A5E", font=cn_font(40, True))
    left, right, top, bottom = 170, 1510, 150, 680
    group_width = (right - left) / 2
    for count in [0, 5, 10, 15]:
        y = bottom - count / 18 * (bottom - top)
        draw.line((left, y, right, y), fill="#E4E7EC", width=2)
        draw.text((105, y - 13), str(count), fill="#667085", font=cn_font(20))
    for dataset_idx, dataset in enumerate(["bci2a", "bci2b"]):
        group_center = left + (dataset_idx + 0.5) * group_width
        for route_idx, route in enumerate(routes):
            row = selection[
                (selection["dataset"] == dataset) & (selection["global_selected"] == route)
            ].iloc[0]
            count = int(row["count"])
            x = group_center + (route_idx - 1.5) * 105
            y = bottom - count / 18 * (bottom - top)
            draw.rectangle((x - 38, y, x + 38, bottom), fill=colors[route_idx])
            draw.text((x - 11, y - 34), str(count), fill="#344054", font=cn_font(21, True))
        draw.text((group_center - 65, bottom + 30), dataset.upper(), fill="#344054", font=cn_font(25, True))
    for idx, (label, color) in enumerate(zip(labels, colors)):
        x = 190 + idx * 345
        draw.rectangle((x, 755, x + 32, 785), fill=color)
        draw.text((x + 44, 752), label, fill="#344054", font=cn_font(20))
    image.save(path, dpi=(180, 180))


def create_figures(data: dict[str, pd.DataFrame]) -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        "pipeline": FIG_DIR / "global_routing_pipeline_zh.png",
        "subjects": FIG_DIR / "subject_improvement_zh.png",
        "selection": FIG_DIR / "route_selection_counts_zh.png",
    }
    create_pipeline_figure(figures["pipeline"])
    create_subject_figure(data, figures["subjects"])
    create_selection_figure(data, figures["selection"])
    return figures


def set_run_cn(run, size=None, bold=None, color=None, italic=None, font_name="Microsoft YaHei"):
    kwargs = {"name": font_name}
    if size is not None:
        kwargs["size"] = size
    if bold is not None:
        kwargs["bold"] = bold
    if color is not None:
        kwargs["color"] = color
    if italic is not None:
        kwargs["italic"] = italic
    base.set_run_font(run, **kwargs)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def configure_document(doc: Document) -> None:
    base.configure_document(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header.paragraphs[0]
    header.text = SHORT_TITLE
    for run in header.runs:
        set_run_cn(run, size=9, color=base.MID_GRAY)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, before, after in [
        ("Title", 23, 0, 12),
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(level=level)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_cn(run, size=16 if level == 1 else 13, bold=True, color=base.BLUE)


def add_body(doc: Document, text: str, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_cn(paragraph.add_run(text), size=10.5, italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.20
        set_run_cn(paragraph.add_run(item), size=10.5)


def set_table_grid(table, widths: list[float]) -> None:
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")


def add_table(doc: Document, title: str, rows: list[list[str]], widths: list[float]) -> None:
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    set_run_cn(caption.add_run(title), size=9.5, bold=True, color=base.DARK_BLUE)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    base.set_table_widths(table, widths)
    set_table_grid(table, widths)
    base.set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            base.set_cell_margins(cell, top=80, start=120, bottom=80, end=120)
            if row_index == 0:
                base.set_cell_shading(cell, base.LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index in {0, 1} else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            set_run_cn(
                paragraph.add_run(str(value)),
                size=7.6 if len(rows[0]) >= 7 else 8.2,
                bold=row_index == 0,
                color=base.DARK_BLUE if row_index == 0 else base.BLACK,
            )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width=6.5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", path.stem)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    set_run_cn(cap.add_run(caption), size=9, color=base.MID_GRAY)


def add_title_page(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(70)
    paragraph.paragraph_format.space_after = Pt(14)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_cn(paragraph.add_run(TITLE), size=23, bold=True, color=base.DARK_BLUE)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(30)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_cn(
        paragraph.add_run("基于 BCI Competition IV-2a/2b 的五随机种子跨会话研究"),
        size=13,
        color=base.MID_GRAY,
    )
    for label, value in [
        ("作者", "[作者姓名]"),
        ("单位", "[院系、机构、城市、国家]"),
        ("通讯作者", "[姓名、邮箱、ORCID]"),
        ("稿件状态", "中文完整初稿，已加入等权消融与稳健性分析"),
        ("代码", "https://github.com/zfqa/eegnetfgl"),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        set_run_cn(paragraph.add_run(f"{label}："), bold=True, color=base.DARK_BLUE)
        set_run_cn(paragraph.add_run(value))
    doc.add_page_break()


def main_tables() -> dict[str, list[list[str]]]:
    return {
        "datasets": [
            ["数据集", "类别", "EEG 通道", "采样率", "协议"],
            ["BCI IV-2a / BNCI2014_001", "4", "22", "250 Hz", "训练会话→训练/验证；评估会话→测试"],
            ["BCI IV-2b / BNCI2014_004", "2", "3", "250 Hz", "训练会话→训练/验证；评估会话→测试"],
        ],
        "main": [
            ["数据集", "方法", "时间窗", "准确率", "Macro-F1", "Kappa"],
            ["2a", "固定 CE", "0-2 s", "0.5672", "0.5550", "0.4229"],
            ["2a", "自适应 CE", "选择的 2 s", "0.5481", "0.5350", "0.3974"],
            ["2a", "固定双教师 KD", "0-2 s", "0.5847", "0.5724", "0.4463"],
            ["2a", "自适应双教师 KD", "选择的 2 s", "0.5353", "0.5224", "0.3805"],
            ["2a", "全局路由", "选择的路线", "0.5931", "0.5842", "0.4575"],
            ["2b", "固定 CE", "0-2 s", "0.7634", "0.7566", "0.5269"],
            ["2b", "自适应 CE", "选择的 2 s", "0.7722", "0.7696", "0.5445"],
            ["2b", "固定双教师 KD", "0-2 s", "0.7677", "0.7610", "0.5354"],
            ["2b", "自适应双教师 KD", "选择的 2 s", "0.7744", "0.7703", "0.5488"],
            ["2b", "全局路由", "选择的路线", "0.7737", "0.7694", "0.5473"],
        ],
        "base_ablation": [
            ["数据集", "监督配置", "alpha_E", "alpha_R", "准确率", "相对 CE"],
            ["2a", "CE", "0", "0", "0.5672", "--"],
            ["2a", "神经单教师", "0.4", "0", "0.5742", "+0.0069"],
            ["2a", "黎曼单教师", "0", "0.4", "0.5655", "-0.0017"],
            ["2a", "双教师", "0.4", "0.4", "0.5847", "+0.0175"],
            ["2b", "CE", "0", "0", "0.7635", "--"],
            ["2b", "神经单教师", "0.4", "0", "0.7667", "+0.0032"],
            ["2b", "黎曼单教师", "0", "0.4", "0.7633", "-0.0002"],
            ["2b", "双教师", "0.4", "0.4", "0.7677", "+0.0042"],
        ],
        "base_ablation_tests": [
            ["数据集", "比较", "差值", "胜/平/负", "dz", "t 检验 p", "Wilcoxon p"],
            ["2a", "双教师 vs CE", "+0.0175", "7/1/1", "1.07", "0.0122", "0.0156"],
            ["2a", "双教师 vs 神经", "+0.0106", "7/1/1", "1.01", "0.0163", "0.0391"],
            ["2a", "双教师 vs 黎曼", "+0.0192", "7/1/1", "0.76", "0.0510", "0.0547"],
            ["2a", "神经 vs CE", "+0.0069", "5/0/4", "0.32", "0.3603", "0.5703"],
            ["2a", "黎曼 vs CE", "-0.0017", "6/1/2", "-0.08", "0.8057", "0.7422"],
            ["2b", "双教师 vs CE", "+0.0042", "6/0/3", "0.26", "0.4574", "0.4258"],
            ["2b", "双教师 vs 神经", "+0.0010", "7/0/2", "0.07", "0.8472", "0.1523"],
            ["2b", "双教师 vs 黎曼", "+0.0044", "7/0/2", "0.29", "0.4173", "0.2031"],
            ["2b", "神经 vs CE", "+0.0032", "6/0/3", "0.29", "0.4060", "0.3789"],
            ["2b", "黎曼 vs CE", "-0.0002", "5/2/2", "-0.03", "0.9338", "0.9375"],
        ],
        "matched": [
            ["数据集", "配置", "总 KD 权重", "准确率", "Macro-F1", "Kappa"],
            ["2a", "双教师 0.2+0.2", "0.4", "0.5852", "0.5732", "0.4469"],
            ["2a", "神经单教师 0.4", "0.4", "0.5742", "0.5620", "0.4322"],
            ["2a", "黎曼单教师 0.4", "0.4", "0.5655", "0.5536", "0.4207"],
            ["2a", "双教师 0.4+0.4", "0.8", "0.5847", "0.5724", "0.4463"],
            ["2a", "神经单教师 0.8", "0.8", "0.5803", "0.5682", "0.4404"],
            ["2a", "黎曼单教师 0.8", "0.8", "0.5689", "0.5555", "0.4252"],
            ["2b", "双教师 0.2+0.2", "0.4", "0.7681", "0.7643", "0.5363"],
            ["2b", "神经单教师 0.4", "0.4", "0.7667", "0.7629", "0.5334"],
            ["2b", "黎曼单教师 0.4", "0.4", "0.7633", "0.7564", "0.5266"],
            ["2b", "双教师 0.4+0.4", "0.8", "0.7677", "0.7610", "0.5354"],
            ["2b", "神经单教师 0.8", "0.8", "0.7669", "0.7623", "0.5338"],
            ["2b", "黎曼单教师 0.8", "0.8", "0.7552", "0.7474", "0.5104"],
        ],
        "matched_tests": [
            ["数据集", "比较", "差值", "胜/平/负", "t 检验 p", "Wilcoxon p"],
            ["2a", "双 0.2+0.2 vs 神经 0.4", "+0.0110", "7/0/2", "0.0567", "0.0547"],
            ["2a", "双 0.2+0.2 vs 黎曼 0.4", "+0.0197", "7/0/2", "0.0192", "0.0273"],
            ["2a", "双 0.4+0.4 vs 神经 0.8", "+0.0044", "7/0/2", "0.3091", "0.2852"],
            ["2a", "双 0.4+0.4 vs 黎曼 0.8", "+0.0158", "5/0/4", "0.2067", "0.4961"],
            ["2b", "双 0.2+0.2 vs 神经 0.4", "+0.0014", "5/1/3", "0.3581", "0.3828"],
            ["2b", "双 0.2+0.2 vs 黎曼 0.4", "+0.0048", "7/0/2", "0.0724", "0.1289"],
            ["2b", "双 0.4+0.4 vs 神经 0.8", "+0.0008", "5/1/3", "0.8574", "0.7422"],
            ["2b", "双 0.4+0.4 vs 黎曼 0.8", "+0.0125", "7/0/2", "0.1934", "0.1367"],
        ],
        "baselines": [
            ["数据集", "方法", "窗口", "准确率", "Macro-F1", "Kappa"],
            ["2a", "全局路由", "选择的 2 s", "0.5931", "0.5842", "0.4575"],
            ["2a", "EEGNetV4 CE", "2 s", "0.5809", "0.5666", "0.4412"],
            ["2a", "黎曼 TS-LogReg", "2 s", "0.5694", "0.5570", "0.4259"],
            ["2a", "CSP-LDA", "2 s", "0.5530", "0.5270", "0.4040"],
            ["2a", "Compact EEGNet CE", "4 s", "0.5836", "0.5749", "0.4448"],
            ["2a", "EEGNetV4 CE", "4 s", "0.6025", "0.5894", "0.4700"],
            ["2a", "黎曼 TS-LogReg", "4 s", "0.6525", "0.6420", "0.5367"],
            ["2a", "CSP-LDA", "4 s", "0.6029", "0.5726", "0.4705"],
            ["2b", "全局路由", "选择的 2 s", "0.7737", "0.7694", "0.5473"],
            ["2b", "EEGNetV4 CE", "2 s", "0.7714", "0.7676", "0.5427"],
            ["2b", "黎曼 TS-LogReg", "2 s", "0.7159", "0.7080", "0.4317"],
            ["2b", "CSP-LDA", "2 s", "0.7136", "0.7076", "0.4272"],
            ["2b", "Compact EEGNet CE", "4 s", "0.7806", "0.7763", "0.5612"],
            ["2b", "EEGNetV4 CE", "4 s", "0.7895", "0.7852", "0.5790"],
            ["2b", "黎曼 TS-LogReg", "4 s", "0.7297", "0.7189", "0.4593"],
            ["2b", "CSP-LDA", "4 s", "0.7303", "0.7197", "0.4606"],
        ],
        "primary_tests": [
            ["数据集", "比较", "准确率差值", "胜/平/负", "t 检验 p", "Wilcoxon p"],
            ["2a", "路由 vs 固定 2 s CE", "+0.0259", "8/0/1", "0.0055", "0.0078"],
            ["2a", "路由 vs 固定双教师", "+0.0084", "6/0/3", "0.1741", "0.2031"],
            ["2a", "路由 vs Compact 4 s CE", "+0.0096", "6/0/3", "0.6466", "0.6523"],
            ["2b", "路由 vs 固定 2 s CE", "+0.0102", "7/0/2", "0.0720", "0.0703"],
            ["2b", "路由 vs 固定双教师", "+0.0059", "6/0/3", "0.4175", "0.3594"],
            ["2b", "路由 vs Compact 4 s CE", "-0.0069", "3/0/6", "0.4160", "0.4961"],
        ],
        "external_2a": [
            ["对比方法", "差值", "胜/平/负", "t 检验 p（Holm）", "Wilcoxon p（Holm）", "dz"],
            ["EEGNetV4 2 s CE", "+0.0123", "7/0/2", "0.1981 (0.5943)", "0.2031 (0.8125)", "0.47"],
            ["黎曼 TS-LogReg 2 s", "+0.0237", "6/0/3", "0.1443 (0.5772)", "0.2383 (0.8125)", "0.54"],
            ["CSP-LDA 2 s", "+0.0401", "9/0/0", "0.0069 (0.0411)", "0.0039 (0.0234)", "1.20"],
            ["CSP-LDA 4 s", "-0.0097", "5/0/4", "0.7148 (1.0000)", "1.0000 (1.0000)", "-0.13"],
            ["EEGNetV4 4 s CE", "-0.0093", "5/0/4", "0.5901 (1.0000)", "1.0000 (1.0000)", "-0.19"],
            ["黎曼 TS-LogReg 4 s", "-0.0594", "1/0/8", "0.0108 (0.0542)", "0.0195 (0.0977)", "-1.10"],
        ],
        "external_2b": [
            ["对比方法", "差值", "胜/平/负", "t 检验 p（Holm）", "Wilcoxon p（Holm）", "dz"],
            ["EEGNetV4 2 s CE", "+0.0023", "4/1/4", "0.4885 (0.4885)", "0.5781 (0.5781)", "0.24"],
            ["黎曼 TS-LogReg 2 s", "+0.0578", "8/0/1", "0.0228 (0.0756)", "0.0195 (0.0781)", "0.94"],
            ["CSP-LDA 2 s", "+0.0601", "7/0/2", "0.0189 (0.0756)", "0.0195 (0.0781)", "0.98"],
            ["CSP-LDA 4 s", "+0.0433", "8/0/1", "0.0120 (0.0600)", "0.0078 (0.0469)", "1.08"],
            ["EEGNetV4 4 s CE", "-0.0158", "2/0/7", "0.1324 (0.2648)", "0.2031 (0.4062)", "-0.56"],
            ["黎曼 TS-LogReg 4 s", "+0.0440", "8/0/1", "0.0058 (0.0345)", "0.0117 (0.0586)", "1.24"],
        ],
        "selection": [
            ["数据集", "固定 CE", "固定双教师", "自适应 CE", "自适应双教师", "固定/自适应"],
            ["BCI IV-2a", "14", "15", "9", "7", "29/16"],
            ["BCI IV-2b", "17", "11", "7", "10", "28/17"],
        ],
        "consensus": [
            ["数据集", "跨 seed 共识", "固定 CE", "原全局路由", "共识 vs CE p", "Wilcoxon p"],
            ["2a", "0.5711", "0.5672", "0.5931", "0.7127", "0.7422"],
            ["2b", "0.7704", "0.7634", "0.7737", "0.1787", "0.4375"],
        ],
        "complexity": [
            ["数据集", "部署路线", "参数量", "MACs", "CPU ms/trial", "GPU ms/trial"],
            ["2a", "固定 2 s Compact", "2,420", "5.89 M", "1.12", "1.01"],
            ["2a", "检测器+Compact", "3,526", "11.25 M", "3.88", "2.36"],
            ["2b", "固定 2 s Compact", "1,634", "0.86 M", "1.48", "1.42"],
            ["2b", "检测器+Compact", "3,500", "4.83 M", "5.31", "3.21"],
        ],
        "classes": [
            ["数据集", "类别", "CE F1", "双教师 F1", "变化"],
            ["2a", "左手", "0.593", "0.610", "+0.018"],
            ["2a", "右手", "0.598", "0.626", "+0.027"],
            ["2a", "双脚", "0.468", "0.492", "+0.024"],
            ["2a", "舌部", "0.595", "0.599", "+0.004"],
            ["2b", "左手", "0.782", "0.782", "+0.000"],
            ["2b", "右手", "0.749", "0.758", "+0.009"],
        ],
    }


def build_docx(figures: dict[str, Path], tables: dict[str, list[list[str]]]) -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    add_heading(doc, "摘要", 1)
    add_body(doc, ABSTRACT)
    paragraph = doc.add_paragraph()
    set_run_cn(paragraph.add_run("关键词："), bold=True, color=base.DARK_BLUE)
    set_run_cn(paragraph.add_run(KEYWORDS))

    add_heading(doc, "1 引言", 1)
    for text in INTRODUCTION:
        add_body(doc, text)
    add_body(doc, "本文主要贡献如下：", italic=True)
    add_bullets(doc, CONTRIBUTIONS)

    add_heading(doc, "2 材料与方法", 1)
    for heading, paragraphs in METHODS.items():
        add_heading(doc, heading, 2)
        for text in paragraphs:
            add_body(doc, text)
        if heading == "2.1 数据集与实验协议":
            add_table(doc, "表 1 数据集特征与跨会话协议", tables["datasets"], [1.75, 0.5, 0.7, 0.65, 2.95])
        elif heading == "2.5 验证集控制全局路由":
            add_figure(
                doc,
                figures["pipeline"],
                "图 1 验证集控制全局路由。训练四个完整的 2 s 候选，仅依据训练会话验证准确率选择一条路线，之后才在评估会话上测试。",
            )

    add_heading(doc, "3 结果", 1)
    for heading, paragraphs in RESULTS.items():
        add_heading(doc, heading, 2)
        for text in paragraphs:
            add_body(doc, text)
        if heading == "3.1 BCI Competition IV-2a 主结果":
            add_table(doc, "表 2 五随机种子候选策略与路由表现", tables["main"], [0.55, 1.55, 1.15, 0.9, 0.9, 0.9])
            add_figure(
                doc,
                figures["subjects"],
                "图 2 全局路由相对固定 2 s CE 的被试平均准确率变化。每名被试先对 5 个随机种子取均值。",
            )
        elif heading == "3.3 神经教师与黎曼教师基础消融":
            add_table(doc, "表 3a 基础教师消融", tables["base_ablation"], [0.55, 1.45, 0.7, 0.7, 0.9, 0.9])
            add_table(doc, "表 3b 基础教师消融配对检验", tables["base_ablation_tests"], [0.45, 1.45, 0.65, 0.75, 0.45, 0.8, 0.9])
        elif heading == "3.4 等总 KD 权重消融":
            add_table(doc, "表 4 等总 KD 权重性能比较", tables["matched"], [0.55, 1.55, 0.8, 0.9, 0.9, 0.9])
            add_table(doc, "表 5 等总 KD 权重配对检验", tables["matched_tests"], [0.5, 2.0, 0.7, 0.8, 0.9, 1.0])
        elif heading == "3.5 与神经及经典基线的比较":
            add_table(doc, "表 6 路由主比较的配对检验", tables["primary_tests"], [0.55, 2.0, 0.9, 0.8, 0.9, 1.0])
            add_table(doc, "表 7 五随机种子神经与经典基线", tables["baselines"], [0.5, 1.75, 0.85, 0.9, 0.9, 0.9])
            add_table(doc, "表 8a 2a 上路由与外部基线的配对比较（括号为 Holm 校正值）", tables["external_2a"], [1.45, 0.65, 0.75, 1.4, 1.55, 0.45])
            add_table(doc, "表 8b 2b 上路由与外部基线的配对比较（括号为 Holm 校正值）", tables["external_2b"], [1.45, 0.65, 0.75, 1.4, 1.55, 0.45])
        elif heading == "3.6 路线使用与被试异质性":
            add_table(doc, "表 9 45 次运行中的验证集路线选择次数", tables["selection"], [0.9, 0.85, 1.0, 0.9, 1.1, 1.25])
            add_figure(doc, figures["selection"], "图 3 两个数据集中各条验证集选择路线的频次。", width=6.3)
        elif heading == "3.7 跨 seed 路由稳健性":
            add_table(doc, "表 10 跨 seed 共识路由稳健性", tables["consensus"], [0.65, 1.1, 0.9, 1.1, 1.1, 1.1])
        elif heading == "3.8 推理复杂度与逐类别表现":
            add_table(doc, "表 11 推理复杂度与 batch-1 延迟", tables["complexity"], [0.55, 1.55, 0.85, 0.85, 1.0, 1.0])
            add_table(doc, "表 12 汇总混淆矩阵得到的逐类别 F1", tables["classes"], [0.65, 1.0, 1.0, 1.0, 0.9])

    add_heading(doc, "4 讨论", 1)
    for text in DISCUSSION:
        add_body(doc, text)
    add_heading(doc, "4.1 局限性", 2)
    add_bullets(doc, LIMITATIONS)

    add_heading(doc, "5 结论", 1)
    add_body(doc, CONCLUSION)

    add_heading(doc, "声明", 1)
    declarations = [
        ("基金资助", "[填写基金项目；如无：本研究未获得专项资助。]"),
        ("利益冲突", "作者声明不存在利益冲突。"),
        ("伦理审批", "本研究仅使用公开、去标识化数据集，未新增受试者数据采集。原始数据的伦理审批与知情同意见原数据集论文。"),
        ("数据可用性", "BCI Competition IV-2a 和 2b 可通过 BNCI Horizon 与 MOABB 获取。"),
        ("代码可用性", "代码仓库：https://github.com/zfqa/eegnetfgl"),
        ("作者贡献", "[按 CRediT 规范填写作者贡献。]"),
        ("致谢", "[如适用，填写致谢。]"),
    ]
    for label, value in declarations:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        set_run_cn(paragraph.add_run(f"{label}："), bold=True, color=base.DARK_BLUE)
        set_run_cn(paragraph.add_run(value))

    add_heading(doc, "参考文献", 1)
    add_body(doc, "参考文献保留原始英文书目信息，以避免期刊名、题名和 DOI 翻译造成检索歧义。", italic=True)
    for index, reference in enumerate(english.REFERENCES, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"{index}. {reference}")
        base.set_run_font(run, name="Arial", size=8.8)

    doc.save(DOCX_PATH)


def markdown_table(rows: list[list[str]]) -> list[str]:
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return lines


def build_markdown(tables: dict[str, list[list[str]]]) -> None:
    lines = [
        f"# {TITLE}",
        "",
        "**作者：** [作者姓名]",
        "**单位：** [院系、机构、城市、国家]",
        "**通讯作者：** [姓名和邮箱]",
        "",
        "## 摘要",
        "",
        ABSTRACT,
        "",
        f"**关键词：** {KEYWORDS}",
        "",
        "## 1 引言",
        "",
    ]
    for text in INTRODUCTION:
        lines.extend([text, ""])
    lines.extend(["本文主要贡献如下：", ""])
    lines.extend([f"- {item}" for item in CONTRIBUTIONS])
    lines.extend(["", "## 2 材料与方法", ""])
    for heading, paragraphs in METHODS.items():
        lines.extend([f"### {heading}", ""])
        for text in paragraphs:
            lines.extend([text, ""])
        if heading == "2.1 数据集与实验协议":
            lines.extend(["**表 1 数据集特征与跨会话协议**", ""])
            lines.extend(markdown_table(tables["datasets"]))
            lines.append("")
    lines.extend(["## 3 结果", ""])
    table_map = {
        "3.1 BCI Competition IV-2a 主结果": [("表 2 五随机种子候选策略与路由表现", "main")],
        "3.3 神经教师与黎曼教师基础消融": [("表 3a 基础教师消融", "base_ablation"), ("表 3b 基础教师消融配对检验", "base_ablation_tests")],
        "3.4 等总 KD 权重消融": [("表 4 等总 KD 权重性能比较", "matched"), ("表 5 等总 KD 权重配对检验", "matched_tests")],
        "3.5 与神经及经典基线的比较": [("表 6 路由主比较的配对检验", "primary_tests"), ("表 7 五随机种子神经与经典基线", "baselines"), ("表 8a 2a 外部基线配对比较", "external_2a"), ("表 8b 2b 外部基线配对比较", "external_2b")],
        "3.6 路线使用与被试异质性": [("表 9 验证集路线选择次数", "selection")],
        "3.7 跨 seed 路由稳健性": [("表 10 跨 seed 共识路由稳健性", "consensus")],
        "3.8 推理复杂度与逐类别表现": [("表 11 推理复杂度与延迟", "complexity"), ("表 12 逐类别 F1", "classes")],
    }
    for heading, paragraphs in RESULTS.items():
        lines.extend([f"### {heading}", ""])
        for text in paragraphs:
            lines.extend([text, ""])
        for title, key in table_map.get(heading, []):
            lines.extend([f"**{title}**", ""])
            lines.extend(markdown_table(tables[key]))
            lines.append("")
    lines.extend(["## 4 讨论", ""])
    for text in DISCUSSION:
        lines.extend([text, ""])
    lines.extend(["### 4.1 局限性", ""])
    lines.extend([f"- {item}" for item in LIMITATIONS])
    lines.extend(["", "## 5 结论", "", CONCLUSION, "", "## 声明", ""])
    lines.extend(
        [
            "**基金资助：** [填写基金声明。]",
            "",
            "**利益冲突：** 作者声明不存在利益冲突。",
            "",
            "**伦理审批：** 本研究使用公开、去标识化数据集。",
            "",
            "**代码可用性：** https://github.com/zfqa/eegnetfgl",
            "",
            "## 参考文献",
            "",
            "参考文献保留原始英文书目信息。",
            "",
        ]
    )
    lines.extend(f"{index}. {reference}" for index, reference in enumerate(english.REFERENCES, 1))
    lines.append("")
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    data = load_data()
    figures = create_figures(data)
    tables = main_tables()
    build_docx(figures, tables)
    build_markdown(tables)
    print(DOCX_PATH)
    print(MD_PATH)


if __name__ == "__main__":
    main()
