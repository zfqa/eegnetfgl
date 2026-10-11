from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent / "referenced_chinese"
SOURCE = ROOT / "中文论文初稿_投稿前详细修改版_期刊引用校正版_FGL引用版_含时间窗消融.docx"
OUTPUT = ROOT / "中文论文初稿_含时间窗消融_补齐基线与复杂度.docx"


def main() -> None:
    document = Document(SOURCE)
    baseline = document.tables[2]
    additions = {
        "CSP-LDA": ("0.5270", "0.4040"),
        "Riemann TS-LogReg": ("0.5570", "0.4259"),
        "EEGNetV4 CE": ("0.5666", "0.4412"),
    }
    for row in baseline.rows[1:]:
        if row.cells[0].text == "2a" and row.cells[1].text in additions:
            row.cells[3].text, row.cells[4].text = additions[row.cells[1].text]

    caption = document.add_paragraph("表 2a  固定与自适应路线的部署参数量及计算量")
    caption.style = "Caption" if "Caption" in document.styles else "Normal"
    table = document.add_table(rows=1, cols=5)
    table.style = baseline.style
    for cell, value in zip(table.rows[0].cells, ["数据集", "部署路线", "输入窗", "参数量", "MMAC/trial"]):
        cell.text = value
    rows = [
        ("2a", "固定 Compact", "2 s", "2,420", "5.89"),
        ("2a", "固定 Compact", "4 s", "3,444", "11.76"),
        ("2a", "检测器 + 2 s Compact", "2 s", "3,526", "11.25"),
        ("2b", "固定 Compact", "2 s", "1,634", "0.86"),
        ("2b", "固定 Compact", "4 s", "2,146", "1.72"),
        ("2b", "检测器 + 2 s Compact", "2 s", "3,500", "4.83"),
    ]
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value

    note = document.add_paragraph(
        "注：参数量由相同 Compact EEGNet 结构在 250 Hz、2 s（500 点）和 4 s（1000 点）输入下统计；"
        "MMAC 为单 trial 前向传播估算值。自适应路线计入检测器，4 s 神经/黎曼教师仅在训练期使用，"
        "不计入部署参数量。模型计算量不等同于从提示出现至判决的总时间。"
    )
    for run in note.runs:
        run.font.size = Pt(8)

    baseline._tbl.addnext(caption._p)
    caption._p.addnext(table._tbl)
    table._tbl.addnext(note._p)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
