"""Insert the completed BCI IV-2b Riemann/spatial ablation into the Chinese DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


DOCX_PATH = (
    Path(__file__).resolve().parent
    / "global_routing_paper"
    / "validation_controlled_global_routing_mi_eeg_manuscript_zh_latest.docx"
)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    new_paragraph = Paragraph(element, paragraph._parent)
    new_paragraph.style = paragraph.style
    new_paragraph.add_run(text)
    return new_paragraph


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise RuntimeError(f"Could not find paragraph beginning with: {prefix}")


def main() -> None:
    doc = Document(DOCX_PATH)
    if any(paragraph.text.startswith("表 4b BCI IV-2b") for paragraph in doc.paragraphs):
        raise RuntimeError("Table 4b already exists; refusing to insert a duplicate.")

    method_paragraph = find_paragraph(doc, "为区分双教师互补")
    method_paragraph.add_run(
        " 在 2b 上，另以 Riemann KD 权重 lambda_R 与空间结构损失权重 lambda_S 的和匹配总辅助权重，"
        "比较联合与单项监督；每个配置均使用 9 名被试和 5 个随机种子。"
    )

    result_anchor = find_paragraph(doc, "2b 的一致口径比较中")
    result_text = (
        "为区分 Riemann KD 与空间结构约束的贡献，在 2b 的 8-32 Hz 协议下补充了等总辅助权重消融（表 4b）。"
        "每个配置均覆盖 9 名被试和 5 个随机种子；先对每名被试的 5 个 seed 求均值，再报告 9 名被试的均值和标准差。"
        "总辅助权重为 0.2 时，联合配置 Riemann KD + spatial (0.1+0.1) 的准确率为 76.48%，"
        "高于仅 Riemann KD (0.2+0) 的 76.08% 和仅 spatial (0+0.2) 的 76.14%，"
        "对应增益分别为 0.40 和 0.34 个百分点。较低总权重的联合配置 (0.05+0.05) 为 75.89%。"
        "该结果表明两项黎曼监督在中等权重下具有互补趋势；由于本组尚未进行与 CE 的配对显著性检验，"
        "该表仅作为描述性消融证据，不作显著性主张。"
    )
    insert_paragraph_after(result_anchor, result_text)

    # Existing table index 5 is "表 4 等总 KD 权重性能比较".
    anchor_table = doc.tables[5]
    caption = doc.add_paragraph("表 4b BCI IV-2b 中 Riemann KD 与空间结构损失的等总辅助权重消融（8-32 Hz）")
    caption.style = find_paragraph(doc, "表 4 等总 KD 权重性能比较").style

    rows = [
        ["配置 (lambda_R + lambda_S)", "总辅助权重", "Acc", "Macro-F1", "Kappa"],
        ["Riemann KD + spatial\n(0.05 + 0.05)", "0.10", "0.7589 +/- 0.1420", "0.7525 +/- 0.1494", "0.5177 +/- 0.2840"],
        ["Riemann KD + spatial\n(0.10 + 0.10)", "0.20", "0.7648 +/- 0.1354", "0.7617 +/- 0.1373", "0.5296 +/- 0.2709"],
        ["Riemann KD only\n(0.20 + 0.00)", "0.20", "0.7608 +/- 0.1373", "0.7543 +/- 0.1458", "0.5216 +/- 0.2745"],
        ["Spatial only\n(0.00 + 0.20)", "0.20", "0.7614 +/- 0.1421", "0.7551 +/- 0.1499", "0.5229 +/- 0.2841"],
    ]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = anchor_table.style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.25), Inches(0.85), Inches(1.02), Inches(1.02), Inches(1.02)]
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.width = widths[column_index]
            set_cell_text(cell, text, bold=row_index == 0 or row_index == 2)

    note = doc.add_paragraph(
        "注：数值为每名被试先对 5 个 seed 求均值后，再跨 9 名被试计算的均值 +/- 标准差；"
        "该标准差描述被试间异质性，不能与按 seed 汇总的标准差直接比较。"
    )
    note.runs[0].font.size = Pt(9)

    # Move newly created blocks directly after Table 4 and before Table 5.
    anchor_table._element.addnext(caption._p)
    caption._p.addnext(table._element)
    table._element.addnext(note._p)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
