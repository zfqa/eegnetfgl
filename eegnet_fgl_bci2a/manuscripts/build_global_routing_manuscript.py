from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_final_manuscript as base


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "manuscripts" / "global_routing_paper"
FIG_DIR = OUT_DIR / "figures"
DOCX_PATH = OUT_DIR / "validation_controlled_global_routing_mi_eeg_manuscript.docx"
MD_PATH = OUT_DIR / "validation_controlled_global_routing_mi_eeg_manuscript.md"

ROUTING_ROOT = ROOT / "results" / "global_routing_multiseed"
PUBLICATION_ROOT = ROOT / "results" / "publication_complete_analysis"

TITLE = (
    "Validation-Controlled Global Routing with Neural and Riemannian Teachers "
    "for Short-Window Motor Imagery EEG Decoding"
)
SHORT_TITLE = "Validation-controlled routing for short-window MI-EEG"

ABSTRACT = (
    "Short-window motor imagery electroencephalography (MI-EEG) decoding can reduce the "
    "amount of signal required before classification, but a single temporal window and a "
    "single supervision strategy do not benefit all subjects equally. This study proposes "
    "validation-controlled global routing, a subject-level strategy-selection framework "
    "that chooses among four 2-s Compact EEGNet candidates: fixed-window cross-entropy "
    "(CE), detector-selected adaptive-window CE, fixed-window dual-teacher knowledge "
    "distillation (KD), and adaptive-window dual-teacher KD. The dual teachers comprise a "
    "frozen 4-s neural EEGNet teacher (Compact EEGNet for dataset 2a and EEGNetV4 for "
    "dataset 2b) and a 4-s Riemannian tangent-space logistic-regression "
    "classifier. Candidate selection uses only training-session validation accuracy; the "
    "evaluation session remains isolated for final testing. Subject-dependent experiments "
    "were conducted on BCI Competition IV datasets 2a and 2b using nine subjects and five "
    "random seeds. On dataset 2a, routing improved accuracy from 56.72% for fixed 2-s CE to "
    "59.31% (+2.59 percentage points), with improvements in eight of nine subjects "
    "(paired t-test p=0.0055; Wilcoxon p=0.0078). The routed result was numerically higher "
    "than Compact EEGNet 4-s CE (59.31% versus 58.36%) but not significantly different "
    "(p=0.6466). On dataset 2b, routing increased fixed 2-s CE from 76.34% to 77.37%, but "
    "the gain was not significant (p=0.0720), and Compact EEGNet 4-s CE remained higher "
    "(78.06%). Across routed runs, the mean effective signal endpoint was 2.11 s on dataset "
    "2a and 2.20 s on dataset 2b, compared with 4 s for the long-window reference. These "
    "results show that validation-controlled routing can accommodate subject-specific "
    "differences in temporal-window and teacher utility, with the strongest evidence on the "
    "four-class dataset 2a. The method should be interpreted as a short-window strategy "
    "selection framework rather than a universal replacement for full-window decoding."
)

KEYWORDS = (
    "brain-computer interface; electroencephalography; motor imagery; temporal routing; "
    "knowledge distillation; Riemannian geometry; EEGNet"
)

INTRODUCTION = [
    (
        "Motor imagery brain-computer interfaces (BCIs) infer imagined movements from "
        "electroencephalography (EEG) and can support communication, rehabilitation, and "
        "device control. Many offline decoding studies use several seconds of post-cue EEG. "
        "Longer windows provide more temporal evidence, but they also increase the amount of "
        "signal that must be observed before a decision can be made. Short-window decoding "
        "is therefore attractive for responsive systems, although early segments may contain "
        "incomplete event-related desynchronization and synchronization patterns."
    ),
    (
        "Compact convolutional models and covariance-based Riemannian classifiers offer "
        "complementary approaches to MI-EEG. EEGNet learns temporal filters and depthwise "
        "spatial filters directly from the waveform, whereas Riemannian methods represent "
        "each trial by a symmetric positive-definite channel covariance matrix and perform "
        "classification with manifold-aware distances or tangent-space features. A compact "
        "neural model is attractive at deployment, while a full-window covariance classifier "
        "can provide robust geometric information during training."
    ),
    (
        "Knowledge distillation allows privileged long-window information to supervise a "
        "short-window student. However, our preliminary experiments showed substantial "
        "subject dependence. A 4-s neural or Riemannian teacher improved some subjects but "
        "caused negative transfer in others. Similarly, detector-selected temporal windows "
        "were useful for some subjects but inferior to the fixed 0-2 s segment for others. "
        "These observations suggest that forcing one temporal or supervisory strategy on "
        "every subject is poorly matched to MI-EEG heterogeneity."
    ),
    (
        "We therefore formulate short-window decoding as validation-controlled strategy "
        "routing. Four complete 2-s candidate decoders are trained for each subject and seed: "
        "fixed-window CE, adaptive-window CE, fixed-window dual-teacher KD, and "
        "adaptive-window dual-teacher KD. A global router selects one candidate using only "
        "training-session validation accuracy. The router is global in the sense that it "
        "chooses one decoding strategy for a subject and training seed; it is not a per-trial "
        "mixture-of-experts gate. At test time, teachers are absent, and the temporal detector "
        "is used only when an adaptive candidate was selected."
    ),
]

CONTRIBUTIONS = [
    (
        "A validation-controlled global routing framework that jointly adapts temporal-window "
        "selection and the use of long-window teacher knowledge."
    ),
    (
        "A four-candidate design combining fixed or detector-selected 2-s windows with CE or "
        "dual-teacher supervision from dataset-specific 4-s neural EEGNet and Riemannian "
        "teachers."
    ),
    (
        "A leakage-controlled subject-dependent evaluation on BCI Competition IV datasets "
        "2a and 2b, covering nine subjects and five random seeds."
    ),
    (
        "Subject-level statistical, route-frequency, long-window, and failure-mode analyses "
        "that distinguish routing error from short-window information limitations."
    ),
]

METHODS = {
    "2.1 Datasets and protocol": [
        (
            "BCI Competition IV dataset 2a was accessed through MOABB as BNCI2014_001. It "
            "contains 22-channel EEG from nine subjects performing four MI classes: left "
            "hand, right hand, both feet, and tongue. Dataset 2b was accessed as BNCI2014_004 "
            "and contains left- versus right-hand MI from nine subjects using three bipolar "
            "EEG channels. Both datasets were sampled at 250 Hz."
        ),
        (
            "Experiments were subject dependent and cross session. Designated training "
            "sessions were divided into stratified training and validation subsets using an "
            "80:20 ratio. Evaluation sessions were retained as untouched test data. This "
            "split was repeated with seeds 42, 2024, 3407, 7, and 123. All model fitting, "
            "detector fitting, checkpoint selection, teacher fitting, and route selection "
            "used training-session data only."
        ),
    ],
    "2.2 Preprocessing and leakage control": [
        (
            "EEG trials were band-pass filtered from 8 to 32 Hz and cropped to 0-4 s after "
            "the cue. Per-channel standardization parameters were estimated from the training "
            "subset and applied unchanged to validation and test data. Fixed short-window "
            "models used the first 2 s. Adaptive windows always retained a 2-s duration but "
            "used a detector-selected starting point."
        ),
        (
            "For dual-teacher training, the 2-s student crop and 4-s teacher input were taken "
            "from the same trial before batching. Labels and trial order were preserved. "
            "Neither teacher was queried during test inference. The test session was not used "
            "to choose a window, detector threshold, model checkpoint, or route."
        ),
    ],
    "2.3 Compact student and long-window teachers": [
        (
            "All routed candidate classifiers used Compact EEGNet with F1=8 temporal filters, "
            "depth multiplier D=2, F2=16 pointwise filters, a 64-sample temporal kernel, ELU "
            "activations, average pooling, and dropout of 0.5. The frozen 4-s neural teacher "
            "was Compact EEGNet for dataset 2a and EEGNetV4 for dataset 2b; each teacher was "
            "trained on the same subject. The Riemannian teacher estimated "
            "4-s trial covariance matrices with Oracle Approximating Shrinkage, projected "
            "them to the tangent space using the Riemannian metric, and fitted multinomial "
            "logistic regression. Its regularization parameter was selected on validation "
            "data from {0.01, 0.03, 0.1, 0.3, 1, 3, 10}."
        ),
        (
            "The dual-teacher objective was L = CE + alpha_E KL_E + alpha_R KL_R. Here CE "
            "is cross-entropy with the true label, KL_E transfers the 4-s neural EEGNet class "
            "distribution, and KL_R transfers the 4-s Riemannian class distribution. The "
            "main setting used alpha_E=0.4, alpha_R=0.4, and temperature T=1.0."
        ),
    ],
    "2.4 Adaptive temporal-window candidates": [
        (
            "Dataset 2a used a class-aware MI-active detector with 0.5-s windows and a 0.25-s "
            "step. Early windows were assigned a non-MI label, stable MI windows were assigned "
            "the trial class, and transition windows were excluded. Detection required a "
            "subject-specific confidence threshold and a specified number of consecutive "
            "positive windows. Threshold and consecutive-window rules had been selected from "
            "training-session validation data in the initial seed-42 pilot and were fixed "
            "before the five-seed routing experiment."
        ),
        (
            "Dataset 2b used an EEGNetV4 detector trained to classify 1-s sliding windows with "
            "a 0.25-s step over the 0-2 s detection range. Confidence scores were smoothed "
            "over three neighboring windows, and the peak-confidence start was used to crop "
            "the final 2-s segment. The dataset-specific detector implementations reflect "
            "differences in channel count and task structure and are not claimed to be a "
            "single unified onset-detection algorithm."
        ),
    ],
    "2.5 Validation-controlled global routing": [
        (
            "For each subject and seed, four complete candidates were trained: fixed CE, "
            "adaptive CE, fixed dual KD, and adaptive dual KD. Every candidate was evaluated "
            "on the same validation subset. The route with the highest validation accuracy "
            "was selected. Exact ties used a conservative priority order: fixed CE, fixed "
            "dual KD, adaptive CE, and adaptive dual KD. After selection, only the chosen "
            "candidate was scored on the evaluation session."
        ),
        (
            "The route is selected once for a subject and training seed rather than for each "
            "test trial. Consequently, a fixed route deploys only Compact EEGNet. An adaptive "
            "route deploys the temporal detector followed by Compact EEGNet. Teachers remain "
            "training-only components in both cases."
        ),
    ],
    "2.6 Optimization and evaluation": [
        (
            "Neural candidates were trained for up to 300 epochs with batch size 32, learning "
            "rate 0.001, weight decay 0.0001, and early-stopping patience of 100 epochs. "
            "Accuracy, macro-F1, and Cohen's kappa were reported. Results were first averaged "
            "over the five seeds within each subject. Paired t-tests and Wilcoxon signed-rank "
            "tests were then applied across the nine subject means. The primary comparison "
            "was global routing versus fixed 2-s Compact EEGNet CE."
        ),
        (
            "Contextual baselines were evaluated under the same subject-dependent, "
            "cross-session, five-seed protocol. They comprised EEGNetV4 trained with CE, "
            "Riemannian tangent-space logistic regression using OAS covariance estimation, "
            "and CSP followed by shrinkage LDA, each evaluated with 2-s and 4-s inputs. "
            "Every reported baseline therefore contains 45 runs (nine subjects multiplied "
            "by five seeds). ShallowConvNet and DeepConvNet were not evaluated and are not "
            "included in the reported comparisons. For comparisons between routing and the "
            "six external baselines, five-seed results were first averaged within each "
            "subject, followed by paired t-tests and Wilcoxon signed-rank tests across nine "
            "subjects. Holm correction was applied separately within each dataset across "
            "the six comparisons."
        ),
        (
            "For practical interpretation, the effective signal endpoint was defined as "
            "2.0 s for a fixed route and adaptive start + 2.0 s for an adaptive route. This "
            "quantity describes the latest EEG sample used by the selected classifier; it "
            "does not include hardware, preprocessing, or neural-network execution time."
        ),
    ],
}

RESULTS = {
    "3.1 Main result on BCI Competition IV-2a": [
        (
            "Global routing achieved 59.31% accuracy, 58.42% macro-F1, and kappa 0.4575. "
            "Fixed 2-s CE achieved 56.72%, 55.50%, and 0.4229, respectively. The routed "
            "accuracy gain was +2.59 percentage points; eight of nine subjects improved and "
            "one declined. The difference was significant by paired t-test (p=0.0055) and "
            "Wilcoxon test (p=0.0078). If the two dataset-level primary comparisons are "
            "Holm-corrected together, the dataset 2a result remains significant "
            "(adjusted p=0.0110 for the t-test)."
        ),
        (
            "The fixed dual-teacher candidate reached 58.47%, confirming that privileged "
            "4-s supervision was useful on average. Routing added a further 0.84 percentage "
            "points, but the routed-versus-fixed-dual difference was not significant "
            "(p=0.1741). The benefit of routing therefore arose from avoiding uniformly poor "
            "candidate choices and exploiting subject-specific variation, rather than from "
            "a large improvement over the strongest single candidate."
        ),
    ],
    "3.2 Cross-dataset result on BCI Competition IV-2b": [
        (
            "On dataset 2b, global routing reached 77.37% accuracy, 76.94% macro-F1, and "
            "kappa 0.5473. Fixed 2-s CE reached 76.34%, 75.66%, and 0.5269. The +1.02-point "
            "accuracy gain favored routing in seven of nine subjects but did not reach "
            "significance (paired t-test p=0.0720; Wilcoxon p=0.0703)."
        ),
        (
            "Adaptive dual KD had a slightly higher unconditionally averaged accuracy than "
            "the selected route (77.44% versus 77.37%). This illustrates validation-to-test "
            "selection noise: a route can be optimal on the available validation subset yet "
            "fail to maximize aggregate test accuracy. The external dataset therefore "
            "supports directional generalization but not a statistically reliable benefit."
        ),
    ],
    "3.3 Comparison with neural and classical baselines": [
        (
            "On dataset 2a, global routing was numerically higher than Compact EEGNet 4-s CE "
            "(59.31% versus 58.36%; +0.96 percentage points; six wins and three losses), but "
            "the difference was not significant (paired t-test p=0.6466; Wilcoxon p=0.6523). "
            "On dataset 2b, Compact EEGNet 4-s CE remained higher (78.06% versus 77.37%; "
            "-0.69 percentage points; paired t-test p=0.4160). Global routing also remained "
            "below EEGNetV4 4-s CE on both datasets."
        ),
        (
            "Among the additional five-seed baselines on dataset 2a, global routing "
            "outperformed EEGNetV4 2-s CE (59.31% versus 58.09%), Riemannian TS-LogReg 2 s "
            "(56.94%), and CSP-LDA 2 s (55.30%). The strongest overall baseline was "
            "Riemannian TS-LogReg 4 s at 65.25%, followed by CSP-LDA 4 s at 60.29% and "
            "EEGNetV4 4-s CE at 60.25%. Thus, routing was competitive among short-window "
            "methods but did not exceed the strongest full-window covariance classifier."
        ),
        (
            "The central strong-backbone comparison did not reach significance: routing "
            "exceeded EEGNetV4 2-s CE by 1.23 percentage points, with seven subject-level "
            "wins and two losses (paired t-test p=0.1981; Wilcoxon p=0.2031). Routing also "
            "exceeded Riemannian TS-LogReg 2 s by 2.37 points without significance "
            "(p=0.1443), whereas its 4.01-point advantage over CSP-LDA 2 s was significant "
            "and survived Holm correction (t-test adjusted p=0.0411; Wilcoxon adjusted "
            "p=0.0234). Compared with 4-s baselines, routing was not different from CSP-LDA "
            "or EEGNetV4 and was lower than Riemannian TS-LogReg; the latter difference was "
            "significant before, but not after, six-comparison Holm correction."
        ),
        (
            "On dataset 2b, EEGNetV4 2-s CE reached 77.14%, slightly below global routing "
            "at 77.37%, whereas Riemannian TS-LogReg 2 s and CSP-LDA 2 s reached 71.59% and "
            "71.36%, respectively. EEGNetV4 4-s CE remained the strongest evaluated method "
            "at 78.95%. CSP-LDA 4 s and Riemannian TS-LogReg 4 s reached 73.03% and 72.97%."
        ),
        (
            "On dataset 2b, the 0.23-point difference from EEGNetV4 2-s CE was not "
            "significant (paired t-test p=0.4885; Wilcoxon p=0.5781). Advantages over the "
            "2-s Riemannian and CSP-LDA baselines were significant before correction but "
            "not after Holm correction. Routing exceeded CSP-LDA 4 s with Holm-adjusted "
            "Wilcoxon p=0.0469 and exceeded Riemannian TS-LogReg 4 s with Holm-adjusted "
            "t-test p=0.0345; the alternative paired test did not remain significant in "
            "either case. Routing remained below EEGNetV4 4-s CE without a significant "
            "difference (p=0.1324)."
        ),
        (
            "The mean effective signal endpoint of the routed strategy was 2.11 s on dataset "
            "2a and 2.20 s on dataset 2b because fixed routes were selected in most runs. "
            "These values are substantially shorter than the 4-s reference, although they "
            "should not be interpreted as measured online decision latency."
        ),
    ],
    "3.4 Route usage and subject heterogeneity": [
        (
            "Across 45 runs per dataset, fixed routes were selected in 64.4% of dataset 2a "
            "runs and 62.2% of dataset 2b runs. Dual-teacher routes were selected in 48.9% "
            "and 46.7%, respectively. No single candidate dominated, supporting the premise "
            "that both temporal-window and teacher utility were subject and seed dependent."
        ),
        (
            "Exploratory test-oracle diagnostics were used only to understand failure modes. "
            "Within the evaluated candidate set, the best 2-s route still remained below "
            "4-s CE for dataset 2a subjects 1, 2, and 3 and dataset 2b subject 9, suggesting "
            "a short-window information limitation. In contrast, dataset 2b subjects 1, 7, "
            "and 8 had a 2-s candidate capable of matching or exceeding 4-s CE, but validation "
            "routing did not consistently select it. These oracle observations were not used "
            "for model selection or headline performance."
        ),
    ],
}

DISCUSSION = [
    (
        "The main finding is that validation-controlled routing significantly improved "
        "Compact EEGNet short-window decoding on the four-class dataset 2a. The result was "
        "not produced by one universally superior expert. Fixed dual KD was the strongest "
        "individual candidate on dataset 2a, while adaptive candidates helped selected "
        "subjects. The router combined these complementary behaviors and avoided applying "
        "temporal adaptation or teacher supervision when validation evidence did not support "
        "them."
    ),
    (
        "The result also clarifies the role of long-window teachers. Distillation was useful "
        "for several difficult subjects, but it was not uniformly beneficial. A 4-s teacher "
        "observes temporal content that may be absent from a 2-s student crop, and the "
        "semantic correspondence becomes weaker when the student window is shifted by the "
        "detector. Validation routing treats teacher knowledge as optional rather than "
        "mandatory, reducing but not eliminating negative transfer."
    ),
    (
        "The comparison with 4-s CE supports a practical but cautious interpretation. On "
        "dataset 2a, the routed 2-s strategy was numerically above Compact EEGNet 4-s CE, "
        "whereas on dataset 2b it was slightly below. Neither comparison was significant. "
        "Accordingly, the evidence supports performance retention with a shorter signal "
        "segment, not general superiority over long-window decoding. A formal non-inferiority "
        "analysis with a prospectively specified margin would strengthen this claim."
    ),
    (
        "Cross-dataset behavior reveals an important limitation. Dataset 2b has only two "
        "classes and three EEG channels, and the fixed 2-s baseline was already high. The "
        "smaller opportunity for improvement, together with validation selection noise, "
        "resulted in a non-significant gain. Moreover, the dataset 2a and 2b detectors use "
        "different objectives. Future work should evaluate a unified detector and an "
        "out-of-fold routing protocol before describing the system as a general temporal "
        "routing architecture."
    ),
    (
        "The broader baseline matrix places the contribution more precisely. Global routing "
        "was higher than the evaluated 2-s EEGNetV4, Riemannian TS-LogReg, and CSP-LDA "
        "baselines on both datasets. However, full-window performance depended strongly on "
        "the dataset: Riemannian TS-LogReg 4 s was strongest on dataset 2a, whereas EEGNetV4 "
        "4 s was strongest on dataset 2b. The proposed contribution is therefore improved "
        "short-window strategy selection, not state-of-the-art accuracy across unrestricted "
        "window lengths."
    ),
    (
        "The strongest next methodological step is out-of-fold route estimation within the "
        "training session. Candidate utility could be estimated across several internal "
        "folds, with a conservative safety threshold that activates adaptive or teacher-guided "
        "routes only when their advantage is stable. Validation-controlled spatial guidance "
        "is another candidate extension, but the current spatial-gating pilot was developed "
        "on one dataset 2b subject and is therefore not included as a main paper result."
    ),
]

LIMITATIONS = [
    (
        "The router selects among separately trained candidates and is not an end-to-end or "
        "per-trial mixture-of-experts gate."
    ),
    (
        "A single 20% validation split is used for route selection; comparing four candidates "
        "can amplify validation noise."
    ),
    (
        "Dataset-specific detector definitions reduce architectural uniformity across datasets."
    ),
    (
        "The effective signal endpoint is reported, but real-time wall-clock latency and "
        "asynchronous operation were not measured."
    ),
    (
        "The method significantly improved fixed 2-s CE only on dataset 2a and did not "
        "significantly outperform fixed dual KD or 4-s CE."
    ),
    (
        "Subject-dependent training limits conclusions about cross-subject transfer and "
        "calibration-free deployment."
    ),
    (
        "ShallowConvNet and DeepConvNet were not included in the completed experimental "
        "matrix; comparison with these convolutional baselines remains future work."
    ),
]

CONCLUSION = (
    "This study reframed short-window MI-EEG decoding as validation-controlled strategy "
    "routing. For each subject, the proposed framework selected among fixed or adaptive "
    "2-s windows and CE or dual-teacher supervision. Across nine subjects and five seeds, "
    "global routing significantly improved fixed 2-s Compact EEGNet on BCI Competition "
    "IV-2a and produced a smaller, non-significant gain on IV-2b. Performance was close to "
    "Compact EEGNet 4-s CE while using a mean effective signal endpoint near 2.1-2.2 s, but "
    "the method did not significantly exceed the long-window reference. The results support "
    "validation-controlled routing as a promising way to manage subject-specific temporal "
    "and teacher heterogeneity, while motivating out-of-fold selection and unified detector "
    "designs for stronger generalization."
)

REFERENCES = base.REFERENCES


def load_data() -> dict[str, pd.DataFrame]:
    return {
        "runs": pd.read_csv(ROUTING_ROOT / "all_subject_seed_results.csv"),
        "subject_means": pd.read_csv(ROUTING_ROOT / "subject_mean_across_seeds.csv"),
        "tests": pd.read_csv(ROUTING_ROOT / "paired_tests_subject_means.csv"),
        "external_tests": pd.read_csv(
            ROUTING_ROOT / "routing_vs_external_baselines_paired_tests.csv"
        ),
        "selection": pd.read_csv(ROUTING_ROOT / "selection_counts.csv"),
        "aggregate": pd.read_csv(ROUTING_ROOT / "aggregate_method_summary.csv"),
    }


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def create_pipeline_figure(path: Path) -> None:
    width, height = 1800, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(42, True)
    label_font = font(28, True)
    body_font = font(24)
    small_font = font(21)
    draw.text((70, 35), "Validation-controlled global routing", fill="#173A5E", font=title_font)

    def box(x0, y0, x1, y1, fill, title, lines):
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=fill, outline="#344054", width=3)
        draw.text((x0 + 22, y0 + 18), title, fill="#101828", font=label_font)
        y = y0 + 64
        for line in lines:
            draw.text((x0 + 22, y), line, fill="#344054", font=body_font)
            y += 34

    box(60, 150, 350, 350, "#E8F1F8", "Training EEG", ["4-s same-trial signal", "Train/validation only"])
    box(430, 105, 790, 245, "#EAF7F2", "Fixed candidates", ["0-2 s CE", "0-2 s dual-teacher KD"])
    box(430, 300, 790, 470, "#FFF4DE", "Adaptive candidates", ["Detector-selected 2-s crop", "CE or dual-teacher KD"])
    box(875, 180, 1215, 390, "#F2F4F7", "Validation router", ["Compare four candidates", "Select highest val. accuracy", "Fixed CE wins exact ties"])
    box(1300, 180, 1720, 390, "#EDE9FE", "Evaluation-session test", ["One selected 2-s strategy", "Teachers absent at inference", "Detector only if adaptive"])
    box(
        430,
        540,
        790,
        690,
        "#FCE7F3",
        "Training-only teachers",
        ["4-s neural EEGNet", "4-s Riemann TS-LogReg"],
    )

    def arrow(x0, y0, x1, y1, color="#475467"):
        draw.line((x0, y0, x1, y1), fill=color, width=5)
        angle = np.arctan2(y1 - y0, x1 - x0)
        length = 18
        for offset in (2.55, -2.55):
            x = x1 + length * np.cos(angle + offset)
            y = y1 + length * np.sin(angle + offset)
            draw.line((x1, y1, x, y), fill=color, width=5)

    arrow(350, 220, 430, 175)
    arrow(350, 285, 430, 370)
    arrow(790, 175, 875, 245)
    arrow(790, 385, 875, 330)
    arrow(1215, 285, 1300, 285)
    arrow(610, 540, 610, 470, "#9D174D")
    draw.text((65, 610), "Global = one route per subject/seed", fill="#667085", font=small_font)
    image.save(path, dpi=(180, 180))


def create_subject_improvement_figure(data: dict[str, pd.DataFrame], path: Path) -> None:
    frame = data["subject_means"].copy()
    width, height = 1800, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 35), "Subject-mean accuracy change: global routing vs fixed 2-s CE", fill="#173A5E", font=font(40, True))
    plot_left, plot_right = 150, 1730
    plot_top, plot_bottom = 140, 760
    center = (plot_top + plot_bottom) // 2
    max_abs = 0.08
    draw.line((plot_left, center, plot_right, center), fill="#667085", width=3)
    draw.text((35, center - 15), "0", fill="#667085", font=font(20))

    group_width = (plot_right - plot_left) / 9
    bar_width = 46
    for subject in range(1, 10):
        x_center = plot_left + (subject - 0.5) * group_width
        for offset, dataset, color in [(-28, "bci2a", "#2E74B5"), (28, "bci2b", "#D97706")]:
            row = frame[(frame["dataset"] == dataset) & (frame["subject"] == subject)].iloc[0]
            value = float(row["global_test_acc"] - row["fixed_ce_test_acc"])
            y = center - value / max_abs * ((plot_bottom - plot_top) / 2)
            top, bottom = sorted([center, y])
            draw.rectangle((x_center + offset - bar_width / 2, top, x_center + offset + bar_width / 2, bottom), fill=color)
        draw.text((x_center - 14, plot_bottom + 24), f"S{subject}", fill="#344054", font=font(21))

    for value in [-0.08, -0.04, 0.04, 0.08]:
        y = center - value / max_abs * ((plot_bottom - plot_top) / 2)
        draw.line((plot_left, y, plot_right, y), fill="#E4E7EC", width=2)
        draw.text((35, y - 13), f"{value*100:+.0f} pp", fill="#667085", font=font(18))
    draw.rectangle((1260, 70, 1295, 100), fill="#2E74B5")
    draw.text((1310, 67), "BCI IV-2a", fill="#344054", font=font(21))
    draw.rectangle((1480, 70, 1515, 100), fill="#D97706")
    draw.text((1530, 67), "BCI IV-2b", fill="#344054", font=font(21))
    draw.text((80, 835), "Positive bars indicate higher accuracy after validation-controlled routing.", fill="#667085", font=font(22))
    image.save(path, dpi=(180, 180))


def create_route_selection_figure(data: dict[str, pd.DataFrame], path: Path) -> None:
    selection = data["selection"].copy()
    routes = ["fixed_ce", "fixed_dual", "adaptive_ce", "adaptive_dual"]
    labels = ["Fixed CE", "Fixed dual KD", "Adaptive CE", "Adaptive dual KD"]
    colors = ["#2E74B5", "#16A085", "#D97706", "#9B51E0"]
    width, height = 1600, 820
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "Route selections across 45 subject-seed runs", fill="#173A5E", font=font(40, True))
    left, right, top, bottom = 170, 1510, 150, 680
    group_width = (right - left) / 2
    max_count = 18
    for count in [0, 5, 10, 15]:
        y = bottom - count / max_count * (bottom - top)
        draw.line((left, y, right, y), fill="#E4E7EC", width=2)
        draw.text((105, y - 13), str(count), fill="#667085", font=font(20))
    for dataset_idx, dataset in enumerate(["bci2a", "bci2b"]):
        group_center = left + (dataset_idx + 0.5) * group_width
        for route_idx, route in enumerate(routes):
            row = selection[(selection["dataset"] == dataset) & (selection["global_selected"] == route)].iloc[0]
            count = int(row["count"])
            x = group_center + (route_idx - 1.5) * 105
            y = bottom - count / max_count * (bottom - top)
            draw.rectangle((x - 38, y, x + 38, bottom), fill=colors[route_idx])
            draw.text((x - 11, y - 34), str(count), fill="#344054", font=font(21, True))
        draw.text((group_center - 65, bottom + 30), "BCI IV-2a" if dataset == "bci2a" else "BCI IV-2b", fill="#344054", font=font(25, True))
    for idx, (label, color) in enumerate(zip(labels, colors)):
        x = 240 + idx * 330
        draw.rectangle((x, 755, x + 32, 785), fill=color)
        draw.text((x + 44, 752), label, fill="#344054", font=font(20))
    image.save(path, dpi=(180, 180))


def create_figures(data: dict[str, pd.DataFrame]) -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        "pipeline": FIG_DIR / "global_routing_pipeline.png",
        "subjects": FIG_DIR / "subject_improvement.png",
        "selection": FIG_DIR / "route_selection_counts.png",
    }
    create_pipeline_figure(figures["pipeline"])
    create_subject_improvement_figure(data, figures["subjects"])
    create_route_selection_figure(data, figures["selection"])
    return figures


def table_rows(data: dict[str, pd.DataFrame]) -> dict[str, list[list[str]]]:
    runs = data["runs"]

    def metrics(dataset: str, prefix: str) -> tuple[str, str, str]:
        group = runs[runs["dataset"] == dataset]
        return (
            f"{group[f'{prefix}_test_acc'].mean():.4f}",
            f"{group[f'{prefix}_test_f1'].mean():.4f}",
            f"{group[f'{prefix}_test_kappa'].mean():.4f}",
        )

    def global_metrics(dataset: str) -> tuple[str, str, str]:
        group = runs[runs["dataset"] == dataset]
        return (
            f"{group['global_test_acc'].mean():.4f}",
            f"{group['global_test_f1'].mean():.4f}",
            f"{group['global_test_kappa'].mean():.4f}",
        )

    method_rows = [["Dataset", "Method", "Window", "Accuracy", "Macro-F1", "Kappa"]]
    for dataset, label in [("bci2a", "BCI IV-2a"), ("bci2b", "BCI IV-2b")]:
        for method, prefix, window in [
            ("Fixed CE", "fixed_ce", "0-2 s"),
            ("Adaptive CE", "adaptive_ce", "Selected 2 s"),
            ("Fixed dual KD", "fixed_dual", "0-2 s"),
            ("Adaptive dual KD", "adaptive_dual", "Selected 2 s"),
        ]:
            acc, f1, kappa = metrics(dataset, prefix)
            method_rows.append([label, method, window, acc, f1, kappa])
        acc, f1, kappa = global_metrics(dataset)
        method_rows.append([label, "Global routing", "Selected route", acc, f1, kappa])

    stats_rows = [
        ["Dataset", "Comparison", "Delta acc.", "W/T/L", "t-test p", "Wilcoxon p"],
        ["2a", "Routing vs fixed 2-s CE", "+0.0259", "8/0/1", "0.0055", "0.0078"],
        ["2a", "Routing vs fixed dual KD", "+0.0084", "6/0/3", "0.1741", "0.2031"],
        ["2a", "Routing vs Compact 4-s CE", "+0.0096", "6/0/3", "0.6466", "0.6523"],
        ["2b", "Routing vs fixed 2-s CE", "+0.0102", "7/0/2", "0.0720", "0.0703"],
        ["2b", "Routing vs fixed dual KD", "+0.0059", "6/0/3", "0.4175", "0.3594"],
        ["2b", "Routing vs Compact 4-s CE", "-0.0069", "3/0/6", "0.4160", "0.4961"],
    ]

    selection = data["selection"]
    selection_rows = [["Dataset", "Fixed CE", "Fixed dual", "Adaptive CE", "Adaptive dual", "Fixed / adaptive"]]
    for dataset, label in [("bci2a", "BCI IV-2a"), ("bci2b", "BCI IV-2b")]:
        counts = {
            row["global_selected"]: int(row["count"])
            for _, row in selection[selection["dataset"] == dataset].iterrows()
        }
        fixed = counts["fixed_ce"] + counts["fixed_dual"]
        adaptive = counts["adaptive_ce"] + counts["adaptive_dual"]
        selection_rows.append(
            [
                label,
                str(counts["fixed_ce"]),
                str(counts["fixed_dual"]),
                str(counts["adaptive_ce"]),
                str(counts["adaptive_dual"]),
                f"{fixed} / {adaptive}",
            ]
        )

    baseline_rows = [
        ["Dataset", "Method", "Window", "Accuracy", "Macro-F1", "Kappa"],
        ["2a", "Global routing", "Selected 2 s", "0.5931", "0.5842", "0.4575"],
        ["2a", "EEGNetV4 CE", "2 s", "0.5809", "0.5666", "0.4412"],
        ["2a", "Riemann TS-LogReg", "2 s", "0.5694", "0.5570", "0.4259"],
        ["2a", "CSP-LDA", "2 s", "0.5530", "0.5270", "0.4040"],
        ["2a", "Compact EEGNet CE", "4 s", "0.5836", "0.5749", "0.4448"],
        ["2a", "EEGNetV4 CE", "4 s", "0.6025", "0.5894", "0.4700"],
        ["2a", "Riemann TS-LogReg", "4 s", "0.6525", "0.6420", "0.5367"],
        ["2a", "CSP-LDA", "4 s", "0.6029", "0.5726", "0.4705"],
        ["2b", "Global routing", "Selected 2 s", "0.7737", "0.7694", "0.5473"],
        ["2b", "EEGNetV4 CE", "2 s", "0.7714", "0.7676", "0.5427"],
        ["2b", "Riemann TS-LogReg", "2 s", "0.7159", "0.7080", "0.4317"],
        ["2b", "CSP-LDA", "2 s", "0.7136", "0.7076", "0.4272"],
        ["2b", "Compact EEGNet CE", "4 s", "0.7806", "0.7763", "0.5612"],
        ["2b", "EEGNetV4 CE", "4 s", "0.7895", "0.7852", "0.5790"],
        ["2b", "Riemann TS-LogReg", "4 s", "0.7297", "0.7189", "0.4593"],
        ["2b", "CSP-LDA", "4 s", "0.7303", "0.7197", "0.4606"],
    ]
    external_rows = [
        [
            "Dataset",
            "Comparator",
            "Delta acc.",
            "W/T/L",
            "t-test p (Holm)",
            "Wilcoxon p (Holm)",
            "dz",
        ]
    ]
    display_names = {
        "EEGNetV4 2s CE": "EEGNetV4 2-s CE",
        "RiemannPaper-TS-LogReg 2s": "Riemann TS-LogReg 2 s",
        "CSP-LDA 2s": "CSP-LDA 2 s",
        "CSP-LDA 4s": "CSP-LDA 4 s",
        "EEGNetV4 4s CE": "EEGNetV4 4-s CE",
        "RiemannPaper-TS-LogReg 4s": "Riemann TS-LogReg 4 s",
    }
    for _, row in data["external_tests"].iterrows():
        external_rows.append(
            [
                "2a" if row["dataset"] == "bci2a" else "2b",
                display_names[row["method"]],
                f"{row['mean_difference']:+.4f}",
                f"{int(row['wins'])}/{int(row['ties'])}/{int(row['losses'])}",
                f"{row['paired_t_p']:.4f} ({row['paired_t_p_holm']:.4f})",
                f"{row['wilcoxon_p']:.4f} ({row['wilcoxon_p_holm']:.4f})",
                f"{row['cohen_dz']:.2f}",
            ]
        )
    return {
        "datasets": [
            ["Dataset", "Classes", "EEG channels", "Sampling", "Protocol"],
            ["BCI IV-2a / BNCI2014_001", "4", "22", "250 Hz", "Training session -> train/validation; evaluation session -> test"],
            ["BCI IV-2b / BNCI2014_004", "2", "3", "250 Hz", "Training sessions -> train/validation; evaluation sessions -> test"],
        ],
        "methods": method_rows,
        "stats": stats_rows,
        "selection": selection_rows,
        "baselines": baseline_rows,
        "external_stats": external_rows,
    }


def set_table_grid(table, widths: list[float]) -> None:
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)


def add_table(doc: Document, title: str, rows: list[list[str]], widths: list[float]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    base.set_run_font(paragraph.add_run(title), size=10, bold=True, color=base.DARK_BLUE)

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    base.set_table_widths(table, widths)
    set_table_grid(table, widths)
    base.set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            base.set_cell_margins(cell, top=80, start=120, bottom=80, end=120)
            if row_index == 0:
                base.set_cell_shading(cell, base.LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index in {0, 1, len(values) - 1}
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            base.set_run_font(
                paragraph.add_run(str(value)),
                size=8.2,
                bold=(row_index == 0),
                color=(base.DARK_BLUE if row_index == 0 else base.BLACK),
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    base.configure_document(doc)
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = SHORT_TITLE
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        base.set_run_font(run, size=9, color=base.MID_GRAY)
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(6)


def add_title_page(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(70)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(paragraph.add_run(TITLE), size=23, bold=True, color=base.DARK_BLUE)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(30)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(
        paragraph.add_run(
            "A five-seed, cross-session study on BCI Competition IV datasets 2a and 2b"
        ),
        size=13,
        color=base.MID_GRAY,
    )

    for label, value in [
        ("Authors", "[Author name(s)]"),
        ("Affiliations", "[Department, Institution, City, Country]"),
        ("Corresponding author", "[Name and email]"),
        ("Target journal", "Medical & Biological Engineering & Computing"),
        ("Manuscript type", "Original research article"),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        base.set_run_font(paragraph.add_run(f"{label}: "), bold=True, color=base.DARK_BLUE)
        base.set_run_font(paragraph.add_run(value))

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(28)
    paragraph.paragraph_format.left_indent = Inches(0.45)
    paragraph.paragraph_format.right_indent = Inches(0.45)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    base.set_run_font(
        paragraph.add_run(
            "Scope: the central claim is a significant improvement over fixed 2-s Compact "
            "EEGNet on BCI IV-2a through validation-controlled strategy selection. The "
            "manuscript does not claim significant superiority over 4-s decoding or universal "
            "generalization to BCI IV-2b."
        ),
        size=10.5,
        italic=True,
        color=base.MID_GRAY,
    )
    doc.add_page_break()


def add_paragraphs(doc: Document, paragraphs: list[str]) -> None:
    for paragraph in paragraphs:
        base.add_body(doc, paragraph)


def build_docx(data: dict[str, pd.DataFrame], figures: dict[str, Path], tables: dict[str, list[list[str]]]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    base.add_heading(doc, "Abstract", 1)
    base.add_body(doc, ABSTRACT)
    paragraph = doc.add_paragraph()
    base.set_run_font(paragraph.add_run("Keywords: "), bold=True, color=base.DARK_BLUE)
    base.set_run_font(paragraph.add_run(KEYWORDS))

    base.add_heading(doc, "1 Introduction", 1)
    add_paragraphs(doc, INTRODUCTION)
    base.add_body(doc, "The contributions of this work are:", italic=True)
    base.add_bullets(doc, CONTRIBUTIONS)

    base.add_heading(doc, "2 Materials and methods", 1)
    for heading, paragraphs in METHODS.items():
        base.add_heading(doc, heading, 2)
        add_paragraphs(doc, paragraphs)
        if heading == "2.1 Datasets and protocol":
            add_table(
                doc,
                "Table 1. Dataset characteristics and cross-session protocol.",
                tables["datasets"],
                [1.75, 0.55, 0.75, 0.75, 2.70],
            )
        elif heading == "2.5 Validation-controlled global routing":
            base.add_figure(
                doc,
                figures["pipeline"],
                "Fig. 1 Validation-controlled global routing. Four complete 2-s candidate "
                "strategies are trained, one is selected by training-session validation "
                "accuracy, and the evaluation session is scored only after route selection.",
                "Diagram of fixed and adaptive CE and dual-teacher candidates feeding a validation router and one selected evaluation strategy.",
                width=6.4,
            )

    base.add_heading(doc, "3 Results", 1)
    for heading, paragraphs in RESULTS.items():
        base.add_heading(doc, heading, 2)
        add_paragraphs(doc, paragraphs)
        if heading == "3.1 Main result on BCI Competition IV-2a":
            add_table(
                doc,
                "Table 2. Five-seed candidate and routed performance (45 runs per dataset).",
                tables["methods"],
                [1.0, 1.55, 1.15, 0.9, 0.9, 1.0],
            )
            base.add_figure(
                doc,
                figures["subjects"],
                "Fig. 2 Subject-mean accuracy change of global routing relative to fixed "
                "2-s CE. Values were first averaged over five seeds within each subject.",
                "Grouped bar chart showing subject-level routing improvements over fixed 2-second cross-entropy for BCI IV-2a and IV-2b.",
                width=6.35,
            )
        elif heading == "3.3 Comparison with neural and classical baselines":
            add_table(
                doc,
                "Table 3. Subject-mean paired comparisons for global routing.",
                tables["stats"],
                [0.6, 2.1, 0.8, 0.7, 0.85, 1.45],
            )
            add_table(
                doc,
                "Table 4. Five-seed neural and classical baseline comparison (45 runs per method and dataset).",
                tables["baselines"],
                [0.55, 1.85, 0.85, 0.95, 0.95, 0.95],
            )
            add_table(
                doc,
                "Table 5. Subject-mean paired comparisons of global routing with external baselines. "
                "Parentheses contain Holm-adjusted p values within each dataset.",
                tables["external_stats"],
                [0.45, 1.55, 0.65, 0.6, 1.15, 1.35, 0.45],
            )
        elif heading == "3.4 Route usage and subject heterogeneity":
            add_table(
                doc,
                "Table 6. Validation-selected route counts across 45 runs.",
                tables["selection"],
                [1.0, 0.9, 1.0, 1.0, 1.15, 1.45],
            )
            base.add_figure(
                doc,
                figures["selection"],
                "Fig. 3 Frequency of each validation-selected route across subjects and seeds.",
                "Bar chart showing counts for fixed CE, fixed dual KD, adaptive CE, and adaptive dual KD on both datasets.",
                width=6.2,
            )

    base.add_heading(doc, "4 Discussion", 1)
    add_paragraphs(doc, DISCUSSION)
    base.add_heading(doc, "4.1 Limitations", 2)
    base.add_bullets(doc, LIMITATIONS)

    base.add_heading(doc, "5 Conclusion", 1)
    base.add_body(doc, CONCLUSION)

    base.add_heading(doc, "Declarations", 1)
    declarations = [
        ("Funding", "[Insert funding statement or: No funding was received for this work.]"),
        ("Competing interests", "The authors declare no competing interests."),
        (
            "Ethics approval",
            "This secondary analysis used publicly available, de-identified datasets and "
            "collected no new participant data. The original dataset publications should be "
            "cited for ethics approval and consent.",
        ),
        ("Consent to participate", "Not applicable to the present secondary analysis."),
        (
            "Data availability",
            "BCI Competition IV datasets 2a and 2b are publicly available through BNCI Horizon and MOABB.",
        ),
        (
            "Code availability",
            "Code is available at https://github.com/zfqa/eegnetfgl after repository cleanup and release tagging.",
        ),
        ("Author contributions", "[Insert CRediT contribution statement.]"),
        ("Acknowledgements", "[Insert acknowledgements, if applicable.]"),
    ]
    for label, value in declarations:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        base.set_run_font(paragraph.add_run(f"{label}: "), bold=True, color=base.DARK_BLUE)
        base.set_run_font(paragraph.add_run(value))

    base.add_heading(doc, "References", 1)
    for index, reference in enumerate(REFERENCES, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(5)
        base.set_run_font(paragraph.add_run(f"{index}. {reference}"), size=9.5)

    doc.save(DOCX_PATH)


def markdown_table(rows: list[list[str]]) -> list[str]:
    output = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
    output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return output


def build_markdown(tables: dict[str, list[list[str]]]) -> None:
    lines = [
        f"# {TITLE}",
        "",
        "**Authors:** [Author name(s)]",
        "**Affiliations:** [Department, Institution, City, Country]",
        "**Corresponding author:** [Name and email]",
        "",
        "## Abstract",
        "",
        ABSTRACT,
        "",
        f"**Keywords:** {KEYWORDS}",
        "",
        "## 1 Introduction",
        "",
    ]
    for paragraph in INTRODUCTION:
        lines.extend([paragraph, ""])
    lines.extend(["The contributions of this work are:", ""])
    lines.extend([f"- {item}" for item in CONTRIBUTIONS])
    lines.extend(["", "## 2 Materials and methods", ""])
    for heading, paragraphs in METHODS.items():
        lines.extend([f"### {heading}", ""])
        for paragraph in paragraphs:
            lines.extend([paragraph, ""])
        if heading == "2.1 Datasets and protocol":
            lines.extend(["**Table 1. Dataset characteristics and protocol.**", ""])
            lines.extend(markdown_table(tables["datasets"]))
            lines.append("")

    lines.extend(["## 3 Results", ""])
    for heading, paragraphs in RESULTS.items():
        lines.extend([f"### {heading}", ""])
        for paragraph in paragraphs:
            lines.extend([paragraph, ""])
        if heading == "3.1 Main result on BCI Competition IV-2a":
            lines.extend(["**Table 2. Candidate and routed performance.**", ""])
            lines.extend(markdown_table(tables["methods"]))
            lines.append("")
        elif heading == "3.3 Comparison with neural and classical baselines":
            lines.extend(["**Table 3. Paired comparisons.**", ""])
            lines.extend(markdown_table(tables["stats"]))
            lines.extend(["", "**Table 4. Five-seed neural and classical baselines.**", ""])
            lines.extend(markdown_table(tables["baselines"]))
            lines.extend(
                [
                    "",
                    "**Table 5. Paired comparisons with external baselines "
                    "(Holm-adjusted p values in parentheses).**",
                    "",
                ]
            )
            lines.extend(markdown_table(tables["external_stats"]))
            lines.append("")
        elif heading == "3.4 Route usage and subject heterogeneity":
            lines.extend(["**Table 6. Route counts.**", ""])
            lines.extend(markdown_table(tables["selection"]))
            lines.append("")

    lines.extend(["## 4 Discussion", ""])
    for paragraph in DISCUSSION:
        lines.extend([paragraph, ""])
    lines.extend(["### 4.1 Limitations", ""])
    lines.extend([f"- {item}" for item in LIMITATIONS])
    lines.extend(["", "## 5 Conclusion", "", CONCLUSION, "", "## Declarations", ""])
    lines.extend(
        [
            "**Funding:** [Insert funding statement.]",
            "",
            "**Competing interests:** The authors declare no competing interests.",
            "",
            "**Ethics approval:** This study used publicly available, de-identified datasets.",
            "",
            "**Code availability:** https://github.com/zfqa/eegnetfgl",
            "",
            "## References",
            "",
        ]
    )
    for index, reference in enumerate(REFERENCES, 1):
        lines.append(f"{index}. {reference}")
    lines.append("")
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    data = load_data()
    figures = create_figures(data)
    tables = table_rows(data)
    build_docx(data, figures, tables)
    build_markdown(tables)
    print(DOCX_PATH)
    print(MD_PATH)


if __name__ == "__main__":
    main()
